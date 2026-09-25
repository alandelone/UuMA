from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from .knowledge_models import ChunkRecord, DocumentVersionRecord, SourceRecord, SourceType
from .knowledge_service import KnowledgeService
from .safe_web import SafeWebFetcher

MAX_DOWNLOAD_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class ExtractedContent:
    text: str
    media_type: str
    title: str
    metadata: dict[str, Any]


class KnowledgeIngestor:
    def __init__(
        self,
        service: KnowledgeService,
        content_dir: Path,
        *,
        fetcher: SafeWebFetcher | None = None,
    ) -> None:
        self.service = service
        self.content_dir = content_dir.resolve()
        self.content_dir.mkdir(parents=True, exist_ok=True)
        self.fetcher = fetcher or SafeWebFetcher(maximum_bytes=MAX_DOWNLOAD_BYTES)

    def ingest_file(self, path: str | Path, *, actor_id: str) -> dict[str, Any]:
        resolved = Path(path).expanduser().resolve()
        if not resolved.is_file():
            raise FileNotFoundError(resolved)
        raw = resolved.read_bytes()
        extracted = self._extract(raw, resolved.suffix.lower(), resolved.name)
        return self._persist(
            locator=resolved.as_uri(),
            raw=raw,
            extracted=extracted,
            source_type=self._source_type(resolved.suffix.lower()),
            actor_id=actor_id,
            metadata={"local_path": str(resolved)},
        )

    def ingest_web(self, url: str, *, actor_id: str) -> dict[str, Any]:
        fetched = self.fetcher.fetch(
            url,
            allowed_content_types={
                "application/pdf",
                "application/xhtml+xml",
                "text/html",
                "text/markdown",
                "text/plain",
            },
        )
        raw = fetched.body
        content_type = fetched.content_type
        final_url = fetched.final_url
        if content_type in {"text/html", "application/xhtml+xml"}:
            probe = raw[:500_000].decode("utf-8", errors="ignore").casefold()
            blocked_markers = (
                "verify you are human",
                "captcha",
                "sign in to continue",
                "log in to continue",
                "subscribe to continue",
                "enable javascript and cookies to continue",
                "access denied",
            )
            if any(marker in probe for marker in blocked_markers):
                raise PermissionError(
                    "The source requires login, CAPTCHA, subscription, or interactive access."
                )
        suffix = {
            "application/pdf": ".pdf",
            "text/html": ".html",
            "application/xhtml+xml": ".html",
            "text/markdown": ".md",
        }.get(content_type, ".txt")
        extracted = self._extract(raw, suffix, final_url)
        return self._persist(
            locator=final_url,
            raw=raw,
            extracted=extracted,
            source_type=SourceType.WEB,
            actor_id=actor_id,
            metadata={"requested_url": url, "content_type": content_type},
        )

    def ingest_text(
        self,
        *,
        locator: str,
        title: str,
        text: str,
        media_type: str = "text/plain",
        source_type: SourceType = SourceType.OTHER,
        actor_id: str,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        raw = text.encode("utf-8")
        return self._persist(
            locator=locator,
            raw=raw,
            extracted=ExtractedContent(text, media_type, title, {}),
            source_type=source_type,
            actor_id=actor_id,
            metadata=metadata or {},
        )

    def _persist(
        self,
        *,
        locator: str,
        raw: bytes,
        extracted: ExtractedContent,
        source_type: SourceType,
        actor_id: str,
        metadata: dict[str, Any],
    ) -> dict[str, Any]:
        digest = hashlib.sha256(raw).hexdigest()
        content_path = self.content_dir / digest[:2] / digest
        content_path.parent.mkdir(parents=True, exist_ok=True)
        if not content_path.exists():
            content_path.write_bytes(raw)

        previous_document: dict[str, Any] | None = None
        source: dict[str, Any] | None = None
        with self.service.store.connect() as conn:
            row = conn.execute(
                "SELECT record_json FROM sources WHERE locator = ? "
                "ORDER BY updated_sequence DESC LIMIT 1",
                (locator,),
            ).fetchone()
            if row:
                source = json.loads(row["record_json"])
                doc_row = conn.execute(
                    "SELECT record_json FROM documents WHERE source_id = ? "
                    "ORDER BY version DESC LIMIT 1",
                    (source["source_id"],),
                ).fetchone()
                if doc_row:
                    previous_document = json.loads(doc_row["record_json"])
            else:
                duplicate = conn.execute(
                    "SELECT record_json FROM documents WHERE content_digest = ? "
                    "ORDER BY updated_sequence DESC LIMIT 1",
                    (digest,),
                ).fetchone()
                if duplicate:
                    duplicate_document = json.loads(duplicate["record_json"])
                    duplicate_source = self.service._require(
                        "sources", duplicate_document["source_id"], conn
                    )
                    return {
                        "source": duplicate_source,
                        "document": duplicate_document,
                        "chunks": self._chunks_for_document(duplicate_document["document_id"]),
                        "unchanged": True,
                        "duplicate_content": True,
                        "duplicate_locator": locator,
                    }
        if previous_document and previous_document["content_digest"] == digest:
            chunks = self._chunks_for_document(previous_document["document_id"])
            return {
                "source": source,
                "document": previous_document,
                "chunks": chunks,
                "unchanged": True,
            }
        if source is None:
            source = self.service.add_source(
                SourceRecord(
                    locator=locator,
                    title=extracted.title,
                    source_type=source_type,
                    content_digest=digest,
                    metadata=metadata | extracted.metadata,
                ),
                actor_id=actor_id,
            )
        version = 1 if previous_document is None else int(previous_document["version"]) + 1
        document = self.service.add_document(
            DocumentVersionRecord(
                source_id=source["source_id"],
                version=version,
                media_type=extracted.media_type,
                language="und",
                content_digest=digest,
                content_ref=str(content_path),
                previous_document_id=(
                    previous_document["document_id"] if previous_document else None
                ),
                metadata=metadata | extracted.metadata,
            ),
            actor_id=actor_id,
        )
        chunks = []
        for ordinal, (start, end, text) in enumerate(self._chunk_text(extracted.text)):
            chunk = self.service.add_chunk(
                ChunkRecord(
                    document_id=document["document_id"],
                    ordinal=ordinal,
                    text=text,
                    location=f"characters {start}-{end}",
                    start_char=start,
                    end_char=end,
                    content_digest=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                ),
                actor_id=actor_id,
            )
            chunks.append(chunk)
        return {"source": source, "document": document, "chunks": chunks, "unchanged": False}

    def _chunks_for_document(self, document_id: str) -> list[dict[str, Any]]:
        with self.service.store.connect() as conn:
            rows = conn.execute(
                "SELECT record_json FROM chunks WHERE document_id = ? ORDER BY ordinal",
                (document_id,),
            ).fetchall()
        return [json.loads(row["record_json"]) for row in rows]

    @staticmethod
    def _chunk_text(
        text: str, *, maximum: int = 1200, overlap: int = 200
    ) -> list[tuple[int, int, str]]:
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        chunks: list[tuple[int, int, str]] = []
        start = 0
        length = len(normalized)
        while start < length:
            end = min(length, start + maximum)
            if end < length:
                boundary = max(
                    normalized.rfind("\n\n", start + maximum // 2, end),
                    normalized.rfind(". ", start + maximum // 2, end),
                    normalized.rfind("。", start + maximum // 2, end),
                )
                if boundary > start:
                    end = boundary + 1
            content = normalized[start:end].strip()
            if content:
                chunks.append((start, end, content))
            if end >= length:
                break
            start = max(start + 1, end - overlap)
        return chunks

    @staticmethod
    def _source_type(suffix: str) -> SourceType:
        if suffix == ".pdf":
            return SourceType.PAPER
        if suffix in {".md", ".txt", ".html", ".htm", ".docx"}:
            return SourceType.MANUAL
        return SourceType.OTHER

    @staticmethod
    def _extract(raw: bytes, suffix: str, name: str) -> ExtractedContent:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
            return ExtractedContent(
                "\n\n".join(pages),
                "application/pdf",
                Path(name).stem,
                {"page_count": len(pages)},
            )
        if suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(raw))
            text = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
            return ExtractedContent(
                text,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                Path(name).stem,
                {"paragraph_count": len(document.paragraphs)},
            )
        decoded = raw.decode("utf-8", errors="replace")
        if suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(decoded, "html.parser")
            for element in soup(["script", "style", "noscript"]):
                element.decompose()
            title = soup.title.get_text(" ", strip=True) if soup.title else Path(name).stem
            text = re.sub(r"\n{3,}", "\n\n", soup.get_text("\n", strip=True))
            return ExtractedContent(text, "text/html", title, {})
        media_type = "text/markdown" if suffix == ".md" else "text/plain"
        return ExtractedContent(decoded, media_type, Path(name).stem or name, {})

    @staticmethod
    def _validate_public_url(url: str) -> None:
        SafeWebFetcher().validate_public_url(url)
