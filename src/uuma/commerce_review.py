"""Commerce Excel 4-sheet review protocol with whole-line self/others ownership enforcement."""

from __future__ import annotations

import hashlib
import io
import json
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill

from .commerce_store import ALLOWED_CATEGORIES, ALLOWED_OWNERSHIPS, CommerceStore

SCHEMA_VERSION = "1.0"
INSTRUCTIONS_SHEET = "说明与摘要"
ORDERS_SHEET = "订单待核对"
PRICE_COMPARE_SHEET = "比价清单"
RECEIPT_CONFIRM_SHEET = "收货确认"


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _sanitize_string(val: Any) -> str:
    s = str(val or "").strip()
    if s.startswith(("=", "+", "-", "@")):
        return "'" + s
    return s


def generate_review_workbook(
    lines: list[dict[str, Any]],
    *,
    workbook_id: str | None = None,
    task_scope: str = "ALL",
) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    w_id = workbook_id or f"wb_{uuid.uuid4().hex[:12]}"
    now = _now()

    # 1. Sheet 1: Instructions & Summary
    ws_info = wb.create_sheet(title=INSTRUCTIONS_SHEET)
    ws_info.views.sheetView[0].showGridLines = True
    info_headers = ["配置 / 元数据项", "值"]
    ws_info.append(info_headers)
    ws_info.append(["schema_version", SCHEMA_VERSION])
    ws_info.append(["workbook_id", w_id])
    ws_info.append(["generated_at", now])
    ws_info.append(["task_scope", task_scope])
    ws_info.append(["total_lines", len(lines)])
    unreviewed_count = sum(1 for line in lines if not line.get("ownership"))
    ws_info.append(["unreviewed_lines", unreviewed_count])
    ws_info.append(["", ""])
    ws_info.append(["重要规则说明", ""])
    ws_info.append(["1. 用途归属", "整行只支持填写 'self'（自用/实验室）或 'others'（代购/其他）。"])
    ws_info.append(["2. 禁止用途拆分", "不支持按数量或比例拆分用途。严禁在 ownership 单元格填写混合比例或拆分子行。"])
    ws_info.append(["3. 电子类入库前提", "仅用户确认整行 'self' 且 category 为 'electronics' 的行可进入收货预览。"])
    ws_info.append(["4. 防注入安全", "订单号、商品 ID 等长标识均按纯文本写入，不得包含可执行公式。"])

    # Style header
    for cell in ws_info[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # 2. Sheet 2: Pending Orders
    ws_orders = wb.create_sheet(title=ORDERS_SHEET)
    ws_orders.views.sheetView[0].showGridLines = True
    order_headers = [
        "line_id",
        "order_id",
        "product_name",
        "variant_name",
        "quantity",
        "currency",
        "unit_price",
        "line_total",
        "category",
        "ownership",
        "ai_category_suggestion",
        "ai_ownership_suggestion",
        "base_evidence_version",
        "base_review_version",
        "notes",
    ]
    ws_orders.append(order_headers)
    for cell in ws_orders[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    for line in lines:
        row = [
            _sanitize_string(line.get("line_id", "")),
            _sanitize_string(line.get("order_id", "")),
            _sanitize_string(line.get("product_name", "")),
            _sanitize_string(line.get("variant_name", "")),
            float(line.get("quantity", 1.0)),
            _sanitize_string(line.get("currency", "MYR")),
            float(line.get("unit_price", 0.0)) if line.get("unit_price") is not None else "",
            float(line.get("line_total", 0.0)) if line.get("line_total") is not None else "",
            _sanitize_string(line.get("category", "electronics")),
            _sanitize_string(line.get("ownership") or ""),
            _sanitize_string(line.get("ai_category_suggestion", "")),
            _sanitize_string(line.get("ai_ownership_suggestion", "")),
            int(line.get("evidence_version", 1)),
            int(line.get("review_version", 1)),
            "",
        ]
        ws_orders.append(row)
        # Ensure IDs are stored as string
        row_idx = ws_orders.max_row
        ws_orders.cell(row=row_idx, column=1).data_type = "s"
        ws_orders.cell(row=row_idx, column=2).data_type = "s"

    # 3. Sheet 3: Price comparison (Phase 2 marker)
    ws_compare = wb.create_sheet(title=PRICE_COMPARE_SHEET)
    ws_compare.append(["状态", "说明"])
    ws_compare.append(["尚未启用", "Phase 2 采价与比价模块尚未交付，本工作簿不伪造报价。"])
    for cell in ws_compare[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    # 4. Sheet 4: Receipt confirmation
    ws_receipt = wb.create_sheet(title=RECEIPT_CONFIRM_SHEET)
    ws_receipt.append(["line_id", "product_name", "quantity", "category", "ownership", "入库资格"])
    for cell in ws_receipt[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for line in lines:
        owner = line.get("ownership") or ""
        cat = line.get("category") or ""
        eligible = "待核对"
        if owner == "self" and cat == "electronics":
            eligible = "可进入收货预览"
        elif owner == "others":
            eligible = "整行排除入库"
        elif owner == "self":
            eligible = "非电子类暂不入库"

        ws_receipt.append(
            [
                _sanitize_string(line.get("line_id", "")),
                _sanitize_string(line.get("product_name", "")),
                float(line.get("quantity", 1.0)),
                cat,
                owner,
                eligible,
            ]
        )

    # Save to memory
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview_review_workbook(
    file_bytes_or_path: bytes | str | Path,
    store: CommerceStore,
) -> dict[str, Any]:
    if isinstance(file_bytes_or_path, (str, Path)):
        wb = openpyxl.load_workbook(str(file_bytes_or_path), data_only=True)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes_or_path), data_only=True)

    if INSTRUCTIONS_SHEET not in wb.sheetnames or ORDERS_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Invalid workbook format. Missing required sheets: {INSTRUCTIONS_SHEET}, {ORDERS_SHEET}."
        )

    # Read schema version from instructions sheet
    ws_info = wb[INSTRUCTIONS_SHEET]
    schema_ver = None
    for row in ws_info.iter_rows(values_only=True):
        if row and row[0] == "schema_version":
            schema_ver = str(row[1]).strip()
            break
    if schema_ver != SCHEMA_VERSION:
        raise ValueError(
            f"Unsupported schema version: '{schema_ver}'. Expected: '{SCHEMA_VERSION}'."
        )

    ws_orders = wb[ORDERS_SHEET]
    rows = list(ws_orders.iter_rows(values_only=True))
    if not rows or len(rows) < 1:
        raise ValueError("Orders sheet is empty.")

    header = [str(c or "").strip() for c in rows[0]]
    col_idx = {name: idx for idx, name in enumerate(header)}
    required_cols = ["line_id", "category", "ownership", "base_evidence_version", "base_review_version"]
    for rc in required_cols:
        if rc not in col_idx:
            raise ValueError(f"Missing required column in orders sheet: '{rc}'.")

    updates = []
    conflicts = []

    for r_num, row in enumerate(rows[1:], start=2):
        if not row or not any(row):
            continue
        line_id_val = str(row[col_idx["line_id"]] or "").strip()
        if not line_id_val:
            continue
        # Remove any leading apostrophe added to prevent formula injection
        line_id_val = line_id_val.removeprefix("'")

        cat_val = str(row[col_idx["category"]] or "").strip().lower()
        if cat_val not in ALLOWED_CATEGORIES:
            conflicts.append(
                {
                    "row": r_num,
                    "line_id": line_id_val,
                    "error_type": "INVALID_CATEGORY",
                    "message": f"Invalid category '{cat_val}'. Allowed: {sorted(ALLOWED_CATEGORIES)}",
                }
            )
            continue

        raw_owner = str(row[col_idx["ownership"]] or "").strip().lower()
        # Enforce no purpose quantity splits
        if any(char in raw_owner for char in [":", "/", ",", ";"]) or any(char.isdigit() for char in raw_owner):
            raise ValueError(
                f"Row {r_num} ({line_id_val}): Purpose quantity splits are strictly prohibited. "
                f"Ownership must be whole-line 'self' or 'others', got: '{raw_owner}'."
            )

        if raw_owner and raw_owner not in ALLOWED_OWNERSHIPS:
            conflicts.append(
                {
                    "row": r_num,
                    "line_id": line_id_val,
                    "error_type": "INVALID_OWNERSHIP",
                    "message": f"Invalid ownership '{raw_owner}'. Must be 'self' or 'others' or blank.",
                }
            )
            continue

        owner_val = raw_owner if raw_owner else None

        base_ev_ver = int(row[col_idx["base_evidence_version"]] or 1)
        base_rev_ver = int(row[col_idx["base_review_version"]] or 1)

        # Check DB state
        with store.connect() as conn:
            existing = conn.execute(
                "SELECT * FROM commerce_order_lines WHERE line_id = ?",
                (line_id_val,),
            ).fetchone()
            if not existing:
                conflicts.append(
                    {
                        "row": r_num,
                        "line_id": line_id_val,
                        "error_type": "UNKNOWN_LINE_ID",
                        "message": f"Line ID '{line_id_val}' does not exist in commerce_order_lines.",
                    }
                )
                continue

            current_ev_ver = existing["evidence_version"]
            current_rev_ver = existing["review_version"]

            if current_rev_ver > base_rev_ver:
                conflicts.append(
                    {
                        "row": r_num,
                        "line_id": line_id_val,
                        "error_type": "STALE_REVIEW_VERSION",
                        "message": f"Review conflict: DB review_version {current_rev_ver} > base {base_rev_ver}.",
                    }
                )
                continue

            if current_ev_ver > base_ev_ver:
                conflicts.append(
                    {
                        "row": r_num,
                        "line_id": line_id_val,
                        "error_type": "STALE_EVIDENCE_VERSION",
                        "message": f"Evidence conflict: Line evidence updated in DB (v{current_ev_ver} > base v{base_ev_ver}).",
                    }
                )
                continue

            # Check if this change modifies current state
            curr_cat = existing["category"]
            curr_owner = existing["ownership"]
            if curr_cat != cat_val or curr_owner != owner_val:
                updates.append(
                    {
                        "line_id": line_id_val,
                        "old_category": curr_cat,
                        "new_category": cat_val,
                        "old_ownership": curr_owner,
                        "new_ownership": owner_val,
                        "base_evidence_version": base_ev_ver,
                        "base_review_version": base_rev_ver,
                    }
                )

    preview_payload = {
        "updates": updates,
        "conflicts": conflicts,
    }
    payload_hash = hashlib.sha256(
        json.dumps(preview_payload, sort_keys=True, ensure_ascii=False).encode("utf-8")
    ).hexdigest()
    preview_id = f"prev_{uuid.uuid4().hex[:12]}"

    return {
        "preview_id": preview_id,
        "payload_hash": payload_hash,
        "updates": updates,
        "conflicts": conflicts,
        "total_updates": len(updates),
        "total_conflicts": len(conflicts),
    }


def commit_review(
    preview: dict[str, Any],
    store: CommerceStore,
    *,
    commit_id: str | None = None,
) -> dict[str, Any]:
    updates = preview.get("updates", [])
    conflicts = preview.get("conflicts", [])
    if conflicts:
        raise ValueError(f"Cannot commit review with {len(conflicts)} unresolved conflicts.")
    if not updates:
        return {"status": "NO_CHANGES", "lines_updated": 0}

    c_id = commit_id or f"rev_commit_{uuid.uuid4().hex[:12]}"
    now = _now()

    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        for u in updates:
            line_id = u["line_id"]
            new_cat = u["new_category"]
            new_owner = u["new_ownership"]
            base_rev_ver = u["base_review_version"]

            # Double check concurrency
            row = conn.execute(
                "SELECT review_version FROM commerce_order_lines WHERE line_id = ?",
                (line_id,),
            ).fetchone()
            if not row or row["review_version"] != base_rev_ver:
                raise ValueError(
                    f"Concurrent modification on line '{line_id}': expected review_version {base_rev_ver}."
                )

            new_rev_ver = base_rev_ver + 1
            review_status = "reviewed" if new_owner else "unreviewed"
            conn.execute(
                """
                UPDATE commerce_order_lines
                SET category = ?, ownership = ?, review_version = ?,
                    review_status = ?, updated_at = ?
                WHERE line_id = ?
                """,
                (new_cat, new_owner, new_rev_ver, review_status, now, line_id),
            )

    return {
        "commit_id": c_id,
        "status": "COMMITTED",
        "lines_updated": len(updates),
        "committed_at": now,
    }
