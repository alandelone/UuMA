from __future__ import annotations

import io

import openpyxl
import pytest

from uuma.commerce_review import (
    commit_review,
    generate_review_workbook,
    preview_review_workbook,
)
from uuma.commerce_store import CommerceStore


@pytest.fixture
def review_setup(tmp_path):
    db_path = tmp_path / "test_lab_review.db"
    archive_dir = tmp_path / "commerce_archive"
    store = CommerceStore(db_path, archive_dir)

    # Ingest baseline batch
    batch = {
        "batch_id": "b_review_01",
        "task_id": "t_01",
        "platform": "shopee",
        "account_id": "acc_01",
        "orders": [
            {
                "order_id": "order_123",
                "items": [
                    {
                        "line_id": "line_review_001",
                        "product_id": "prod_1",
                        "product_name": "STM32F401 MCU",
                        "quantity": 5.0,
                        "unit_price": 12.0,
                        "line_total": 60.0,
                    },
                    {
                        "line_id": "line_review_002",
                        "product_id": "prod_2",
                        "product_name": "Soldering Flux",
                        "quantity": 2.0,
                        "unit_price": 15.0,
                        "line_total": 30.0,
                    },
                ],
            }
        ],
    }
    store.ingest_batch(batch)
    return store


def test_generate_and_preview_review_workbook(review_setup):
    store = review_setup
    lines = store.list_order_lines()
    assert len(lines) == 2

    # 1. Generate 4-sheet workbook
    wb_bytes = generate_review_workbook(lines)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    assert "说明与摘要" in wb.sheetnames
    assert "订单待核对" in wb.sheetnames
    assert "比价清单" in wb.sheetnames
    assert "收货确认" in wb.sheetnames

    # 2. Modify in-memory workbook: set ownership=self for line 1, others for line 2
    ws = wb["订单待核对"]
    # Row 1 is header, Row 2 is line 1, Row 3 is line 2
    # column 10 is ownership (1-indexed)
    ws.cell(row=2, column=10, value="self")
    ws.cell(row=3, column=10, value="others")

    mod_buf = io.BytesIO()
    wb.save(mod_buf)
    mod_bytes = mod_buf.getvalue()

    # 3. Preview diff
    preview = preview_review_workbook(mod_bytes, store)
    assert preview["total_updates"] == 2
    assert preview["total_conflicts"] == 0

    # 4. Commit review
    commit_res = commit_review(preview, store)
    assert commit_res["status"] == "COMMITTED"
    assert commit_res["lines_updated"] == 2

    # Verify updated DB state
    updated_lines = {l["line_id"]: l for l in store.list_order_lines()}
    assert updated_lines["line_review_001"]["ownership"] == "self"
    assert updated_lines["line_review_001"]["review_status"] == "reviewed"
    assert updated_lines["line_review_002"]["ownership"] == "others"


def test_reject_purpose_quantity_splitting(review_setup):
    store = review_setup
    lines = store.list_order_lines()

    wb_bytes = generate_review_workbook(lines)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["订单待核对"]

    # Attempt quantity split: "self: 3, others: 2"
    ws.cell(row=2, column=10, value="self: 3, others: 2")

    mod_buf = io.BytesIO()
    wb.save(mod_buf)

    with pytest.raises(ValueError, match="Purpose quantity splits are strictly prohibited"):
        preview_review_workbook(mod_buf.getvalue(), store)


def test_review_version_conflict_detection(review_setup):
    store = review_setup
    lines = store.list_order_lines()
    wb_bytes = generate_review_workbook(lines)

    # First user/session commits a review to line 1
    with store.connect() as conn:
        conn.execute(
            "UPDATE commerce_order_lines SET review_version = 5 WHERE line_id = 'line_review_001'"
        )

    # Second user tries to submit workbook with base_review_version = 1
    # Line 1 should have STALE_REVIEW_VERSION conflict if changed
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["订单待核对"]
    ws.cell(row=2, column=10, value="self")
    mod_buf = io.BytesIO()
    wb.save(mod_buf)

    preview2 = preview_review_workbook(mod_buf.getvalue(), store)
    assert preview2["total_conflicts"] > 0
    conflict = preview2["conflicts"][0]
    assert conflict["error_type"] == "STALE_REVIEW_VERSION"
