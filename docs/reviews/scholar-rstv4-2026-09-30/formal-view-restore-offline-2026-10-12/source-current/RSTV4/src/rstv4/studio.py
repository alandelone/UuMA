from __future__ import annotations

import csv
import difflib
import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import threading
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor
from contextlib import nullcontext
from datetime import UTC, datetime
from functools import partial
from pathlib import Path
from typing import Any

from . import studio_response_checks as response_checks
from .database import initialize_catalog, transaction
from .manuscript_scientific_review import (
    scientific_binding_stamp,
    scientific_generation_is_stale,
)
from .model_gateway import GatewayDeferredError, UnlimitedLLMGateway
from .models import SubmissionEventInput, new_id
from .service import NotFoundError, RevisionConflictError, RSTV4Error
from .studio_change_links import list_changes, original_ids, record_initial, set_links
from .studio_context import (
    bound_quote,
    citation_ids,
    structure_object,
)
from .studio_correspondence import (
    has_correspondence,
    realization_document,
)
from .studio_correspondence import (
    unresolved as unresolved_correspondence,
)
from .studio_correspondence import (
    validate as validate_correspondence,
)
from .studio_correspondence import (
    validate_origins as validate_correspondence_origins,
)
from .studio_delivery import file_inventory, list_packages, package_content
from .studio_discussions import verify_projection
from .studio_materials import freeze as freeze_materials
from .studio_materials import latex as material_latex
from .studio_materials import verify_files as verify_material_files
from .studio_read_views import (
    document_revision as read_document_revision,
)
from .studio_read_views import (
    generated_markdown,
    verified_artifact_bytes,
)
from .studio_read_views import (
    generated_version as read_generated_version,
)
from .studio_read_views import (
    search_sources as read_sources,
)
from .studio_realization import OUTPUT_SCHEMA
from .studio_realization import validate as validate_realization
from .studio_revision_mapping import write_mapping
from .studio_rounds import baseline_text, content_digest
from .studio_rounds import resolve as resolve_review_round
from .studio_rounds import version as review_baseline_version
from .studio_sources import freeze_source, model_source, original_pdf
from .studio_structure import (
    argument_context,
    paragraph_group,
    section_levels,
    section_titles,
    validate_argument_relations,
    validate_groups,
)
from .studio_suggestions import MODEL_FORMAT, preview_change
from .studio_suggestions import apply as apply_suggestion
from .studio_suggestions import normalize as normalize_suggestions
from .studio_templates import apply_structure, empty_sections, template_preview
from .track_first import TrackFirstService

WORKFLOW_STAGES = (
    "PLANNING",
    "DRAFTING",
    "WHOLE_PAPER_REVIEW",
    "FINALIZATION_SUBMISSION",
    "REVISION",
    "PUBLICATION",
)

_STATUS_BY_STAGE = {
    "PLANNING": "PLANNING",
    "DRAFTING": "DRAFTING",
    "WHOLE_PAPER_REVIEW": "INTERNAL_REVIEW",
    "FINALIZATION_SUBMISSION": "READY_TO_SUBMIT",
    "REVISION": "REVISING",
    "PUBLICATION": "READY_TO_SUBMIT",
}

_ALLOWED_TRANSITIONS = {
    "PLANNING": {"DRAFTING"},
    "DRAFTING": {"PLANNING", "WHOLE_PAPER_REVIEW"},
    "WHOLE_PAPER_REVIEW": {"DRAFTING", "FINALIZATION_SUBMISSION"},
    "FINALIZATION_SUBMISSION": {"REVISION", "PUBLICATION"},
    "REVISION": {"REVISION", "FINALIZATION_SUBMISSION", "PUBLICATION"},
    "PUBLICATION": {"PUBLICATION"},
}

_ASSET_TYPES = {
    "image/png": "FIGURE",
    "image/jpeg": "FIGURE",
    "image/svg+xml": "FIGURE",
    "text/csv": "TABLE",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "TABLE",
    "text/markdown": "TABLE",
    "application/pdf": "DOCUMENT",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "DOCUMENT",
    "application/x-tex": "DOCUMENT",
    "text/plain": "DOCUMENT",
}


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _validate_review(response: dict[str, Any]) -> None:
    findings = response.get("findings")
    if not isinstance(findings, list):
        raise RSTV4Error("Whole-paper model review must explicitly return a findings array")
    for item in findings:
        if (not isinstance(item, dict) or not isinstance(item.get("detail"), str)
                or not item["detail"].strip() or not isinstance(item.get("kind", "SEMANTIC_CHECK"), str)
                or not item.get("kind", "SEMANTIC_CHECK").strip()
                or not isinstance(item.get("severity", "WARNING"), str)
                or item.get("severity", "WARNING") not in {"INFO", "WARNING", "ERROR"}
                or not isinstance(item.get("anchor", {}), dict)):
            raise RSTV4Error("A whole-paper model finding has invalid detail, kind, severity or anchor")


def _decode(row: Any, *json_fields: str) -> dict[str, Any]:
    result = dict(row)
    for field in json_fields:
        value = result.pop(field)
        result[field.removesuffix("_json")] = json.loads(value) if value else None
    return result


def _escape_latex(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(char, char) for char in text)


def _default_document(manuscript: dict[str, Any]) -> dict[str, Any]:
    section_id = new_id("section")
    paragraph_id = new_id("paragraph")
    intent_id = new_id("intent")
    return {
        "schema_version": 1,
        "manuscript_id": manuscript["manuscript_id"],
        "title": manuscript["title"],
        "abstract": "",
        "venue": "",
        "sections": [
            {
                "id": section_id,
                "title": "Introduction",
                "role": "Motivation",
                "paragraphs": [
                    {
                        "id": paragraph_id,
                        "purpose": "",
                        "paradigm": "",
                        "intents": [
                            {
                                "id": intent_id,
                                "label": "Research Context",
                                "thoughts": [],
                            }
                        ],
                        "asset_ids": [],
                        "formulas": [],
                    }
                ],
            }
        ],
    }


def _legacy_document(manuscript: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "manuscript_id": manuscript["manuscript_id"],
        "title": manuscript["title"],
        "abstract": "",
        "venue": "",
        "sections": [],
    }


class DocumentValidationError(RSTV4Error):
    def __init__(self, issues: list[dict[str, str]]) -> None:
        self.issues = issues
        super().__init__("The manuscript structure is not ready to generate")


class JobCancelled(RSTV4Error):
    pass


class StudioService:
    """Author-controlled PaperSkeleton editor backed by RSTV4's catalog database."""

    def __init__(self, root: str | Path | None = None) -> None:
        self.root = Path(root or Path(__file__).resolve().parents[3]).resolve()
        self.db_path = self.root / "data" / "catalog.db"
        self.storage_root = self.root / "data" / "writing-studio"
        self.storage_root.mkdir(parents=True, exist_ok=True)
        initialize_catalog(self.db_path)
        self.track_service = TrackFirstService(self.root)
        self.gateway = UnlimitedLLMGateway(self.root)
        self._executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="writing-studio")
        self._running: set[str] = set()
        self._run_lock = threading.Lock()
        with transaction(self.db_path) as conn:
            conn.execute(
                "UPDATE studio_jobs SET status='QUEUED',started_at=NULL "
                "WHERE status='RUNNING'"
            )

    def list_tracks(self) -> list[dict[str, Any]]:
        with transaction(self.db_path) as conn:
            tracks = [dict(row) for row in conn.execute(
                "SELECT * FROM research_tracks ORDER BY updated_at DESC"
            )]
            for track in tracks:
                track["manuscripts"] = [dict(row) for row in conn.execute(
                    """SELECT manuscript_id,title,workflow_stage,writing_status,revision,
                       updated_at FROM manuscripts WHERE track_id=? ORDER BY updated_at DESC""",
                    (track["track_id"],),
                )]
        return tracks

    def create_manuscript(
        self,
        track_id: str,
        title: str,
        scope: str | None = None,
    ) -> dict[str, Any]:
        from .models import TrackFirstManuscriptInput

        manuscript = self.track_service.create_manuscript(
            TrackFirstManuscriptInput(track_id=track_id, title=title, scope=scope)
        )
        self._ensure_document(manuscript, seed_new=True)
        return self.get_workspace(manuscript["manuscript_id"])

    def get_workspace(self, manuscript_id: str) -> dict[str, Any]:
        manuscript = self.track_service.get_manuscript(manuscript_id)
        row = self._ensure_document(manuscript)
        with transaction(self.db_path) as conn:
            assets = [self._asset(row) for row in conn.execute(
                "SELECT * FROM studio_assets WHERE manuscript_id=? ORDER BY created_at",
                (manuscript_id,),
            )]
            latest_job = conn.execute(
                "SELECT * FROM studio_jobs WHERE manuscript_id=? ORDER BY created_at DESC LIMIT 1",
                (manuscript_id,),
            ).fetchone()
            latest_artifacts = [self._artifact(item) for item in conn.execute(
                "SELECT * FROM studio_artifacts WHERE manuscript_id=? ORDER BY created_at DESC",
                (manuscript_id,),
            )]
            findings = [self._finding(item) for item in conn.execute(
                "SELECT * FROM studio_findings WHERE manuscript_id=? ORDER BY created_at DESC",
                (manuscript_id,),
            )]
            template_instances = [
                _decode(item, "adopted_json")
                for item in conn.execute(
                    """SELECT * FROM studio_template_instances WHERE manuscript_id=?
                       ORDER BY created_at""",
                    (manuscript_id,),
                )
            ]
            scientific_stale = scientific_generation_is_stale(conn, manuscript_id, row["active_generated_version_id"])
        return {
            "manuscript": manuscript,
            "workflow_stage": manuscript.get("workflow_stage") or "PLANNING",
            "document_revision": row["revision"],
            "document": json.loads(row["document_json"]),
            "active_generated_version_id": row["active_generated_version_id"],
            "last_generated_revision": row["last_generated_revision"],
            "generated_is_stale": (
                scientific_stale or (row["last_generated_revision"] is not None
                and row["last_generated_revision"] != row["revision"])
            ),
            "assets": assets,
            "latest_job": self._job(latest_job) if latest_job else None,
            "artifacts": latest_artifacts,
            "findings": findings,
            "template_instances": template_instances,
        }

    def get_document_revision(self, manuscript_id: str, revision: int) -> dict[str, Any]:
        """Read an actual author revision, preserving legacy editor initialization."""
        self.get_workspace(manuscript_id)
        with transaction(self.db_path) as conn:
            return read_document_revision(conn, manuscript_id, revision)

    def _verified_artifact_bytes(self, artifact: Any) -> bytes:
        return verified_artifact_bytes(self.storage_root, artifact)

    def get_generated_version(self, manuscript_id: str, version_id: str) -> dict[str, Any]:
        self.track_service.get_manuscript(manuscript_id)
        with transaction(self.db_path) as conn:
            return read_generated_version(conn, self.storage_root, manuscript_id, version_id)

    def generated_material_content(self, manuscript_id: str, version_id: str,
                                   asset_id: str, rendition: str) -> tuple[bytes, str, str]:
        if rendition not in {"source", "render"}:
            raise NotFoundError("Unknown generated material representation")
        version = self.get_generated_version(manuscript_id, version_id)
        asset = next((a for block in version["blocks"] for a in block.get("assets", [])
                      if a["asset_id"] == asset_id), None)
        if asset is None:
            raise NotFoundError("Material does not belong to this generated version")
        key = "source_file" if rendition == "source" else "render_file"
        if key not in asset:
            raise NotFoundError("This material has no image representation")
        with transaction(self.db_path) as conn:
            row = conn.execute("SELECT inputs_json FROM manuscript_versions WHERE manuscript_id=? AND manuscript_version_id=?",
                               (manuscript_id, version_id)).fetchone()
            job_id = json.loads(row["inputs_json"])["job_id"]
        directory = self.storage_root / manuscript_id / "generated" / job_id
        verify_material_files(directory, [asset])
        return ((directory / asset[key]).read_bytes(),
                asset["media_type"] if rendition == "source" else "image/png",
                asset["filename"] if rendition == "source" else asset["asset_id"] + ".png")

    def generated_source_content(self, manuscript_id: str, version_id: str, evidence_id: str) -> bytes:
        version = self.get_generated_version(manuscript_id, version_id)
        reference = next((r for r in version["references"] if r["evidence_id"] == evidence_id), None)
        if reference is None:
            raise NotFoundError("Source does not belong to this generated version")
        with transaction(self.db_path) as conn:
            return original_pdf(conn, reference)

    def list_model_calls(self, manuscript_id: str, *, job_id: str | None = None) -> list[dict[str, Any]]:
        self.track_service.get_manuscript(manuscript_id)
        with transaction(self.db_path) as conn:
            if job_id and not conn.execute(
                "SELECT 1 FROM studio_jobs WHERE manuscript_id=? AND job_id=?",
                (manuscript_id, job_id),
            ).fetchone():
                raise NotFoundError("The model call job does not belong to this manuscript")
            rows = conn.execute(
                "SELECT * FROM studio_model_calls WHERE manuscript_id=?"
                + (" AND job_id=?" if job_id else "") + " ORDER BY started_at,call_id",
                (manuscript_id, job_id) if job_id else (manuscript_id,),
            ).fetchall()
        return [{**{key: value for key, value in dict(row).items() if not key.endswith("_json")},
                 **{key[:-5]: json.loads(row[key]) if row[key] else None
                    for key in ["input_json", "output_json", "contract_json", "error_json"]}}
                for row in rows]

    def list_changes(self, manuscript_id: str) -> list[dict[str, Any]]:
        return list_changes(self.db_path, manuscript_id)

    def set_change_links(self, manuscript_id: str, change_id: str, **decision: Any) -> dict[str, Any]:
        return set_links(self.db_path, manuscript_id, change_id, **decision)

    def _model_call(
        self, manuscript_id: str, task: str, payload: dict[str, Any], *, input_revision: int,
        job_id: str | None = None, message_id: str | None = None,
        validate: Callable[[dict[str, Any]], Any] | None = None,
    ) -> tuple[dict[str, Any], str]:
        call_id, started = new_id("studio_model_call"), _now()
        input_json = _json(payload)
        response = None
        status, error = "FAILED", None
        try:
            response = self.gateway.analyze(task, payload)
            if not isinstance(response, dict):
                raise RSTV4Error("Studio model response must be a JSON object")
            status = "REJECTED"
            if validate is not None:
                validate(response)
            status = "VALIDATED" if validate else "RETURNED"
            return response, call_id
        except Exception as exc:
            error = {"type": type(exc).__name__, "message": str(exc)}
            raise
        finally:
            with transaction(self.db_path) as conn:
                conn.execute(
                    "INSERT INTO studio_model_calls VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (call_id, manuscript_id, job_id, message_id, task, input_revision,
                     hashlib.sha256(input_json.encode("utf-8")).hexdigest(), input_json,
                     _json(response) if response is not None else None,
                     _json(response.get("_contract")) if isinstance(response, dict) and response.get("_contract") is not None else None,
                     status, _json(error) if error else None, started, _now()),
                )

    def _chat_anchor(
        self, manuscript_id: str, anchor_type: str, anchor_id: str | None,
        anchor_revision: int, quote: dict[str, Any] | None,
    ) -> dict[str, Any]:
        quote = quote or {}
        if anchor_type.startswith("generated_"):
            source = quote.get("source")
            version_id = quote.get("generated_version_id") or (source if isinstance(source, dict) else {}).get(
                "generated_version_id"
            )
            if not isinstance(version_id, str):
                raise ValueError("Generated discussion requires its original version")
            generated = self.get_generated_version(manuscript_id, version_id)
            block = next((item for item in generated["blocks"] if item["anchor_id"] == anchor_id), None)
            kind = anchor_type.removeprefix("generated_")
            if block is None or block["kind"] != kind:
                raise ValueError("The object does not exist in the requested generated version")
            if generated["input_revision"] is not None and anchor_revision != generated["input_revision"]:
                raise ValueError("Discussion revision does not match the generated version")
            if generated["input_revision"] is None and anchor_revision != 0:
                raise ValueError("The historical generated version has no known author revision")
            binding = bound_quote(
                manuscript_id=manuscript_id, anchor_type=anchor_type, anchor_id=anchor_id,
                anchor_revision=anchor_revision, label=block["label"], text=block["display_text"],
                raw_text=block["text"], quote=quote, generated_version_id=version_id,
            )
            binding["source"]["author_revision_known"] = generated["input_revision"] is not None
            binding["source"]["generated_version"] = generated["version"]
            return {"quote": binding, "generated": generated, "block": block}
        snapshot = self.get_document_revision(manuscript_id, anchor_revision)
        label, text = structure_object(snapshot["document"], anchor_type, anchor_id)
        binding = bound_quote(
            manuscript_id=manuscript_id, anchor_type=anchor_type, anchor_id=anchor_id,
            anchor_revision=anchor_revision, label=label, text=text, raw_text=text, quote=quote,
        )
        return {"quote": binding, "snapshot": snapshot}

    def get_chat_context(self, manuscript_id: str, message_id: str) -> dict[str, Any]:
        with transaction(self.db_path) as conn:
            row = conn.execute(
                "SELECT * FROM studio_chat_messages WHERE manuscript_id=? AND message_id=?",
                (manuscript_id, message_id),
            ).fetchone()
            if row is not None:
                verify_projection(conn, row)
        if row is None:
            raise NotFoundError("The discussion does not belong to this manuscript")
        message = self._message(row)
        context = self._chat_anchor(
            manuscript_id, message["anchor_type"], message["anchor_id"],
            message["anchor_revision"], message["quoted"],
        )
        return {"message_id": message_id, **context}

    def adopt_template(
        self,
        manuscript_id: str,
        *,
        source_skeleton_id: str | None,
        source_version: str | None,
        adopted: dict[str, Any],
    ) -> dict[str, Any]:
        """Retain a legacy metadata record; it never applies a template to author text."""
        self.track_service.get_manuscript(manuscript_id)
        template_instance_id = new_id("template_instance")
        timestamp = _now()
        with transaction(self.db_path) as conn:
            conn.execute(
                "INSERT INTO studio_template_instances VALUES(?,?,?,?,?,?,?)",
                (
                    template_instance_id,
                    manuscript_id,
                    source_skeleton_id,
                    source_version,
                    _json({**adopted, "source_verification": "UNVERIFIED_LEGACY_RECORD"}),
                    timestamp,
                    timestamp,
                ),
            )
            row = conn.execute(
                """SELECT * FROM studio_template_instances
                   WHERE template_instance_id=?""",
                (template_instance_id,),
            ).fetchone()
        return _decode(row, "adopted_json")

    def preview_template(self, manuscript_id: str, skeleton_id: str) -> dict[str, Any]:
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN")
            return template_preview(conn, manuscript_id, skeleton_id)

    def apply_template(self, manuscript_id: str, request: dict[str, Any]) -> dict[str, Any]:
        """Commit one exact, source-checked template adoption with its author revision."""
        if request["mode"] not in {"INSERT", "REPLACE"}:
            raise ValueError("Unknown template adoption mode")
        if request["mode"] == "REPLACE" and request["replace_confirmed"] is not True:
            raise ValueError("Template replacement requires explicit author confirmation")
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            operation = conn.execute(
                "SELECT * FROM studio_operations WHERE operation_id=?",
                (request["operation_id"],),
            ).fetchone()
            if operation:
                instance = conn.execute(
                    """SELECT * FROM studio_template_instances WHERE manuscript_id=?
                       AND json_extract(adopted_json,'$.operation_id')=?""",
                    (manuscript_id, request["operation_id"]),
                ).fetchone()
                if (operation["manuscript_id"] != manuscript_id or instance is None
                        or json.loads(instance["adopted_json"]).get("request") != request):
                    raise RevisionConflictError("Template operation already has another request")
                row = conn.execute(
                    "SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,),
                ).fetchone()
                return {**self._saved_response(row, dict(operation), idempotent=True),
                        "template_instance": _decode(instance, "adopted_json")}
            row = conn.execute(
                "SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,),
            ).fetchone()
            if row is None:
                raise NotFoundError("Unknown Studio document")
            if row["revision"] != request["base_revision"]:
                raise RevisionConflictError("Author revision changed; preview the template again")
            preview = template_preview(conn, manuscript_id, request["skeleton_id"])
            if preview["source_sha256"] != request["source_sha256"]:
                raise RevisionConflictError("Template source changed; inspect its current version")
            before = json.loads(row["document_json"])
            sections = empty_sections(preview, request["selected_unit_ids"])
            document = apply_structure(before, sections, request["scope_type"],
                                       request["scope_id"], request["mode"])
            saved = self.save_document(
                manuscript_id, base_revision=request["base_revision"],
                operation_id=request["operation_id"], document=document,
                operation_kind="ADOPT_TEMPLATE", comment_ids=request["comment_ids"],
                _connection=conn,
            )
            instance_id = new_id("template_instance")
            adopted = {"operation_id": request["operation_id"], "request": request,
                       "source": preview, "resulting_revision": saved["document_revision"]}
            timestamp = _now()
            conn.execute("INSERT INTO studio_template_instances VALUES(?,?,?,?,?,?,?)",
                         (instance_id, manuscript_id, preview["skeleton_id"],
                          str(preview["skeleton_version"]), _json(adopted), timestamp, timestamp))
            instance = conn.execute(
                "SELECT * FROM studio_template_instances WHERE template_instance_id=?",
                (instance_id,),
            ).fetchone()
            return {**saved, "template_instance": _decode(instance, "adopted_json")}

    def resolve_finding(self, finding_id: str, status: str) -> dict[str, Any]:
        if status not in {"OPEN", "ACCEPTED", "DISMISSED", "RESOLVED"}:
            raise ValueError("Unknown finding status")
        with transaction(self.db_path) as conn:
            if conn.execute(
                "SELECT 1 FROM studio_findings WHERE finding_id=?", (finding_id,)
            ).fetchone() is None:
                raise NotFoundError(f"Unknown whole-paper finding: {finding_id}")
            conn.execute(
                "UPDATE studio_findings SET status=?,resolved_at=? WHERE finding_id=?",
                (status, _now() if status != "OPEN" else None, finding_id),
            )
            row = conn.execute(
                "SELECT * FROM studio_findings WHERE finding_id=?", (finding_id,)
            ).fetchone()
        return self._finding(row)

    def save_document(
        self,
        manuscript_id: str,
        *,
        base_revision: int,
        operation_id: str,
        document: dict[str, Any],
        operation_kind: str = "EDIT",
        inverse: dict[str, Any] | None = None,
        comment_ids: list[str] | None = None,
        suggestion_decision: dict[str, Any] | None = None,
        _connection: Any = None,
    ) -> dict[str, Any]:
        self._validate_document_shape(manuscript_id, document)
        with (nullcontext(_connection) if _connection is not None else transaction(self.db_path)) as conn:
            # Serialize the revision read with its writes, including provenance.
            if _connection is None:
                conn.execute("BEGIN IMMEDIATE")
            previous = conn.execute(
                "SELECT * FROM studio_operations WHERE operation_id=?", (operation_id,)
            ).fetchone()
            if previous is not None:
                if previous["manuscript_id"] != manuscript_id:
                    raise RevisionConflictError("Operation id already belongs to another manuscript")
                if (
                    previous["base_revision"] != base_revision
                    or previous["operation_kind"] != operation_kind
                    or json.loads(previous["payload_json"])["document"] != document
                    or json.loads(previous["payload_json"]).get("suggestion_decision") != suggestion_decision
                ):
                    raise RevisionConflictError("Operation id already belongs to a different edit")
                original_comment_ids = original_ids(conn, previous)
                if original_comment_ids != set(comment_ids or []):
                    raise RevisionConflictError("Operation id already has different review links")
                row = conn.execute(
                    "SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,)
                ).fetchone()
                return self._saved_response(row, dict(previous), idempotent=True)

            row = conn.execute(
                "SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,)
            ).fetchone()
            if row is None:
                raise NotFoundError(f"Unknown Studio document: {manuscript_id}")
            if row["revision"] != base_revision:
                raise RevisionConflictError(
                    f"Studio revision changed from {base_revision} to {row['revision']}"
                )
            for comment_id in set(comment_ids or []):
                comment = conn.execute(
                    """SELECT r.manuscript_id,c.archived_at FROM review_comments c
                       JOIN review_rounds r ON r.review_round_id=c.review_round_id
                       WHERE c.comment_id=?""", (comment_id,),
                ).fetchone()
                if comment is None:
                    raise NotFoundError(f"Unknown review comment: {comment_id}")
                if comment["manuscript_id"] != manuscript_id or comment["archived_at"]:
                    raise ValueError("Review comment must be active and belong to this manuscript")
            before = json.loads(row["document_json"])
            changed_block_ids = self._changed_block_ids(before, document)
            validate_correspondence_origins(conn, manuscript_id, document)
            resulting_revision = base_revision + 1
            timestamp = _now()
            conn.execute(
                """UPDATE studio_documents SET revision=?,document_json=?,updated_at=?
                   WHERE manuscript_id=?""",
                (resulting_revision, _json(document), timestamp, manuscript_id),
            )
            if changed_block_ids:
                placeholders = ",".join("?" for _ in changed_block_ids)
                conn.execute(
                    f"DELETE FROM studio_sample_candidates WHERE manuscript_id=? "
                    f"AND anchor_id IN ({placeholders})",
                    (manuscript_id, *sorted(changed_block_ids)),
                )
            conn.execute(
                """INSERT INTO studio_operations VALUES(?,?,?,?,?,?,?,?)""",
                (
                    operation_id,
                    manuscript_id,
                    base_revision,
                    resulting_revision,
                    operation_kind,
                    _json({"document": document, "comment_ids": sorted(set(comment_ids or [])),
                           **({"suggestion_decision": suggestion_decision} if suggestion_decision else {})}),
                    _json(inverse or {"document": before}),
                    timestamp,
                ),
            )
            change_id = new_id("change")
            conn.execute(
                "INSERT INTO studio_changes VALUES(?,?,?,?,?,?,?)",
                (
                    change_id,
                    manuscript_id,
                    operation_id,
                    _json(before),
                    _json(document),
                    1 if operation_kind in {"UNDO", "REDO"} else 0,
                    timestamp,
                ),
            )
            for comment_id in comment_ids or []:
                comment = conn.execute(
                    "SELECT 1 FROM review_comments WHERE comment_id=?", (comment_id,)
                ).fetchone()
                if comment is None:
                    raise NotFoundError(f"Unknown review comment: {comment_id}")
                conn.execute(
                    "INSERT OR IGNORE INTO comment_changes VALUES(?,?)",
                    (comment_id, change_id),
                )
            record_initial(conn, change_id, comment_ids or [], origin="recorded_at_save", timestamp=timestamp)
            conn.execute(
                """UPDATE manuscripts SET workflow_stage=CASE
                   WHEN workflow_stage='PLANNING' THEN 'DRAFTING' ELSE workflow_stage END,
                   writing_status=CASE WHEN workflow_stage='PLANNING' THEN 'DRAFTING'
                   ELSE writing_status END,updated_at=? WHERE manuscript_id=?""",
                (timestamp, manuscript_id),
            )
            updated = conn.execute(
                "SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,)
            ).fetchone()
            if suggestion_decision:
                self._insert_suggestion_decision(conn, manuscript_id, suggestion_decision, operation_id)
        return self._saved_response(
            updated,
            {
                "operation_id": operation_id,
                "operation_kind": operation_kind,
                "resulting_revision": resulting_revision,
                "created_at": timestamp,
            },
            idempotent=False,
        )

    def set_stage(
        self,
        manuscript_id: str,
        stage: str,
        *,
        expected_revision: int,
        reason: str,
    ) -> dict[str, Any]:
        if stage not in WORKFLOW_STAGES:
            raise ValueError(f"Unknown workflow stage: {stage}")
        next_revision = expected_revision + 1
        timestamp = _now()
        with transaction(self.db_path) as conn:
            manuscript = conn.execute(
                "SELECT * FROM manuscripts WHERE manuscript_id=?", (manuscript_id,),
            ).fetchone()
            if manuscript is None:
                raise NotFoundError(f"Unknown manuscript: {manuscript_id}")
            current = manuscript["workflow_stage"] or "PLANNING"
            if stage != current and stage not in _ALLOWED_TRANSITIONS[current]:
                raise ValueError(f"Cannot move manuscript from {current} to {stage}")
            if manuscript["revision"] != expected_revision:
                raise RevisionConflictError("Manuscript revision changed")
            if stage not in {"PLANNING", "DRAFTING"}:
                self._require_successful_version(
                    conn, manuscript_id,
                    require_current=(
                        stage in {"FINALIZATION_SUBMISSION", "PUBLICATION"}
                        or (stage == "WHOLE_PAPER_REVIEW" and current != stage)
                    ),
                )
            if (
                stage in {"FINALIZATION_SUBMISSION", "PUBLICATION"}
                and manuscript["needs_scientific_review"]
            ):
                raise RSTV4Error(
                    "Review the manuscript's scientific dependencies before finalization"
                )
            conn.execute(
                """UPDATE manuscripts SET workflow_stage=?,writing_status=?,revision=?,updated_at=?
                   WHERE manuscript_id=?""",
                (stage, _STATUS_BY_STAGE[stage], next_revision, timestamp, manuscript_id),
            )
            conn.execute(
                """INSERT INTO manuscript_status_history
                   VALUES(?,?,?,?,?,?,?,?)""",
                (
                    new_id("history"),
                    manuscript_id,
                    _STATUS_BY_STAGE[stage],
                    None,
                    None,
                    reason,
                    next_revision,
                    timestamp,
                ),
            )
        return self.get_workspace(manuscript_id)

    @staticmethod
    def _require_successful_version(
        conn: Any, manuscript_id: str, *, require_current: bool,
    ) -> None:
        studio = conn.execute(
            "SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,),
        ).fetchone()
        version = None
        if studio is not None and studio["active_generated_version_id"]:
            version = conn.execute(
                "SELECT * FROM manuscript_versions "
                "WHERE manuscript_version_id=? AND manuscript_id=?",
                (studio["active_generated_version_id"], manuscript_id),
            ).fetchone()
        if version is not None:
            provenance = json.loads(version["inputs_json"])
            job = conn.execute(
                "SELECT * FROM studio_jobs WHERE job_id=? AND manuscript_id=?",
                (provenance.get("job_id"), manuscript_id),
            ).fetchone()
            if (
                job is not None and job["job_kind"] == "GENERATE" and job["status"] == "SUCCEEDED"
                and job["input_revision"] == studio["last_generated_revision"]
                and provenance.get("studio_revision") == studio["last_generated_revision"]
            ):
                if require_current and studio["last_generated_revision"] != studio["revision"]:
                    raise RSTV4Error("Regenerate the manuscript before advancing this stage")
                if require_current and scientific_generation_is_stale(conn, manuscript_id, studio["active_generated_version_id"]):
                    manuscript = conn.execute("SELECT needs_scientific_review FROM manuscripts WHERE manuscript_id=?", (manuscript_id,)).fetchone()
                    if manuscript["needs_scientific_review"]:
                        raise RSTV4Error("Review the manuscript's scientific dependencies before advancing this stage")
                    raise RSTV4Error("Regenerate after the manuscript's scientific binding review")
                return
        raise RSTV4Error("A successful generated version is required before advancing this stage")

    def search_sources(
        self, manuscript_id: str, query: str, *, limit: int = 30,
    ) -> list[dict[str, Any]]:
        self.track_service.get_manuscript(manuscript_id)
        with transaction(self.db_path) as conn:
            return read_sources(conn, manuscript_id, query, limit=limit)

    def get_samples(
        self,
        manuscript_id: str,
        *,
        layer: str,
        anchor_id: str | None = None,
        role: str | None = None,
        intent: str | None = None,
        limit: int = 12,
    ) -> list[dict[str, Any]]:
        manuscript = self.track_service.get_manuscript(manuscript_id)
        studio = self._ensure_document(manuscript)
        actual_anchor = anchor_id or "paper"
        query_key = hashlib.sha256(
            _json({"layer": layer, "role": role, "intent": intent, "limit": limit}).encode()
        ).hexdigest()
        with transaction(self.db_path) as conn:
            cached = conn.execute(
                """SELECT sample_json FROM studio_sample_candidates
                   WHERE manuscript_id=? AND anchor_id=? AND document_revision=? AND query_key=?
                   ORDER BY created_at,candidate_id""",
                (manuscript_id, actual_anchor, studio["revision"], query_key),
            ).fetchall()
        if cached:
            return [json.loads(row["sample_json"]) for row in cached]
        conditions = ["(e.track_id=? OR e.field_id=t.field_id)"]
        parameters: list[Any] = [manuscript["track_id"]]
        if role:
            conditions.append("su.layer4 LIKE ?")
            parameters.append(f"%{role}%")
        if intent:
            conditions.append("su.layer3_json LIKE ?")
            parameters.append(f"%{intent}%")
        parameters.append(min(max(limit, 1), 50))
        with transaction(self.db_path) as conn:
            rows = conn.execute(
                f"""SELECT su.unit_id,su.section,su.raw_text,su.layer1_json,su.layer2_json,
                    su.layer3_json,su.layer4,ps.skeleton_id,p.title,p.paper_id,pv.version
                    FROM skeleton_units su
                    JOIN paper_skeletons ps ON ps.skeleton_id=su.skeleton_id
                    JOIN paper_versions pv ON pv.paper_version_id=ps.paper_version_id
                    JOIN papers p ON p.paper_id=pv.paper_id
                    JOIN evidence e ON e.paper_id=p.paper_id
                    JOIN research_tracks t ON t.track_id=?
                    WHERE {' AND '.join(conditions)}
                    GROUP BY su.unit_id ORDER BY p.published_year DESC,su.seq_idx LIMIT ?""",
                [manuscript["track_id"], *parameters],
            ).fetchall()
        results = [
            {
                **dict(row),
                "layer": layer,
                "layer1": json.loads(row["layer1_json"]),
                "layer2": json.loads(row["layer2_json"]),
                "layer3": json.loads(row["layer3_json"]),
            }
            for row in rows
        ]
        with transaction(self.db_path) as conn:
            for item in results:
                conn.execute(
                    """INSERT OR REPLACE INTO studio_sample_candidates
                       VALUES(?,?,?,?,?,?,?)""",
                    (
                        manuscript_id,
                        actual_anchor,
                        studio["revision"],
                        query_key,
                        item["unit_id"],
                        _json(item),
                        _now(),
                    ),
                )
        return results

    def store_asset(
        self,
        manuscript_id: str,
        filename: str,
        media_type: str | None,
        content: bytes,
    ) -> dict[str, Any]:
        self.track_service.get_manuscript(manuscript_id)
        if not content:
            raise ValueError("Asset is empty")
        if len(content) > 25 * 1024 * 1024:
            raise ValueError("Asset exceeds the 25 MB limit")
        guessed = mimetypes.guess_type(filename)[0]
        actual_type = (media_type or guessed or "application/octet-stream").split(";", 1)[0]
        if actual_type not in _ASSET_TYPES:
            raise ValueError(f"Unsupported asset type: {actual_type}")
        metadata: dict[str, Any] = {}
        if actual_type == "image/svg+xml":
            metadata = self._validate_svg(content)
        elif actual_type == "text/csv":
            metadata = self._preview_csv(content)
        elif actual_type == "text/markdown":
            metadata = self._preview_markdown_table(content)
        elif actual_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            metadata = self._preview_xlsx(content)
        elif actual_type == "application/pdf":
            metadata = self._validate_pdf_asset(content)
        elif actual_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            metadata = self._validate_docx_asset(content)
        elif actual_type in {"application/x-tex", "text/plain"}:
            try:
                preview = content.decode("utf-8")[:1000]
            except UnicodeDecodeError as exc:
                raise ValueError("Text assets must be UTF-8") from exc
            metadata = {"encoding": "utf-8", "preview": preview}
        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name).strip("._")
        if not safe_name:
            safe_name = "asset"
        sha256 = hashlib.sha256(content).hexdigest()
        asset_dir = self.storage_root / manuscript_id / "assets"
        asset_dir.mkdir(parents=True, exist_ok=True)
        destination = asset_dir / f"{sha256[:16]}-{safe_name}"
        if not destination.exists():
            destination.write_bytes(content)
        asset_id = new_id("asset")
        with transaction(self.db_path) as conn:
            existing = conn.execute(
                "SELECT * FROM studio_assets WHERE manuscript_id=? AND sha256=?",
                (manuscript_id, sha256),
            ).fetchone()
            if existing is None:
                conn.execute(
                    """INSERT INTO studio_assets VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        asset_id,
                        manuscript_id,
                        safe_name,
                        actual_type,
                        len(content),
                        sha256,
                        str(destination),
                        _ASSET_TYPES[actual_type],
                        "READY",
                        _json(metadata),
                        _now(),
                    ),
                )
                existing = conn.execute(
                    "SELECT * FROM studio_assets WHERE asset_id=?", (asset_id,)
                ).fetchone()
        return self._asset(existing)

    def list_chat(
        self,
        manuscript_id: str,
        *,
        anchor_id: str | None = None,
        before: str | None = None,
        limit: int = 100,
        anchor_type: str | None = None,
        anchor_revision: int | None = None,
        generated_version_id: str | None = None,
        before_message_id: str | None = None,
    ) -> list[dict[str, Any]]:
        conditions = ["manuscript_id=?"]
        parameters: list[Any] = [manuscript_id]
        if anchor_id:
            conditions.append("anchor_id=?")
            parameters.append(anchor_id)
        for column, value in (("anchor_type", anchor_type), ("anchor_revision", anchor_revision)):
            if value is not None:
                conditions.append(f"{column}=?")
                parameters.append(value)
        if generated_version_id is not None:
            conditions.append(
                "COALESCE(json_extract(quoted_json,'$.source.generated_version_id'),"
                "json_extract(quoted_json,'$.generated_version_id'))=?"
            )
            parameters.append(generated_version_id)
        parameters.append(min(max(limit, 1), 200))
        with transaction(self.db_path) as conn:
            if before_message_id is not None:
                cursor = conn.execute(
                    "SELECT rowid,created_at FROM studio_chat_messages "
                    "WHERE manuscript_id=? AND message_id=?",
                    (manuscript_id, before_message_id),
                ).fetchone()
                if cursor is None:
                    raise NotFoundError("The history cursor does not belong to this manuscript")
                if before is not None and before != cursor["created_at"]:
                    raise ValueError("The history cursor time does not match its original message")
                conditions.append("(created_at<? OR (created_at=? AND rowid<?))")
                parameters[-1:-1] = [cursor["created_at"], cursor["created_at"], cursor["rowid"]]
            elif before:
                conditions.append("created_at<?")
                parameters[-1:-1] = [before]
            rows = conn.execute(
                f"""SELECT * FROM studio_chat_messages WHERE {' AND '.join(conditions)}
                    ORDER BY created_at DESC,rowid DESC LIMIT ?""",
                parameters,
            ).fetchall()
            for row in rows:
                verify_projection(conn, row)
        return [self._message(row) for row in reversed(rows)]

    def _suggestion(self, manuscript_id: str, message_id: str, suggestion_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
        with transaction(self.db_path) as conn:
            row = conn.execute(
                "SELECT * FROM studio_chat_messages WHERE manuscript_id=? AND message_id=? AND role='AGENT'",
                (manuscript_id, message_id),
            ).fetchone()
            if row is not None:
                verify_projection(conn, row)
        if row is None:
            raise NotFoundError("The suggestion message does not belong to this manuscript")
        message = self._message(row)
        suggestion = next((item for item in message["quoted"].get("suggestions", []) if item["suggestion_id"] == suggestion_id), None)
        if suggestion is None:
            raise NotFoundError("The message has no such suggestion")
        return message, suggestion

    def suggestion_preview(self, manuscript_id: str, message_id: str, suggestion_id: str) -> dict[str, Any]:
        message, suggestion = self._suggestion(manuscript_id, message_id, suggestion_id)
        workspace = self.get_workspace(manuscript_id)
        with transaction(self.db_path) as conn:
            changes = [preview_change(workspace["document"], change) for change in suggestion["changes"]]
            for change in changes:
                if change["action"] == "ADD_SOURCE" and not self._suggestion_source_available(conn, manuscript_id, change["evidence"]["evidence_id"]):
                    change.update(available=False, source_unavailable=True)
            decisions = [_decode(row, "request_json", "provenance_json") for row in conn.execute(
                "SELECT * FROM studio_suggestion_decisions WHERE manuscript_id=? AND message_id=? AND suggestion_id=? ORDER BY created_at",
                (manuscript_id, message_id, suggestion_id),
            )]
        adopted = {change_id for decision in decisions if decision["outcome"] != "REJECT"
                   for change_id in decision["request"]["selected_change_ids"]}
        return {"suggestion": suggestion, "source": message["quoted"]["source"],
                "document_revision": workspace["document_revision"],
                "changes": [{**change, "already_adopted": change["change_id"] in adopted} for change in changes],
                "decisions": decisions}

    @staticmethod
    def _suggestion_source_available(conn: Any, manuscript_id: str, evidence_id: str) -> bool:
        return conn.execute(
            """SELECT 1 FROM evidence e JOIN manuscripts m ON m.manuscript_id=?
               JOIN research_tracks t ON t.track_id=m.track_id
               WHERE e.evidence_id=? AND (e.track_id=m.track_id OR e.field_id=t.field_id)""",
            (manuscript_id, evidence_id),
        ).fetchone() is not None

    @staticmethod
    def _insert_suggestion_decision(conn: Any, manuscript_id: str, decision: dict[str, Any], operation_id: str | None) -> None:
        conn.execute("INSERT INTO studio_suggestion_decisions VALUES(?,?,?,?,?,?,?,?,?)", (
            decision["decision_id"], manuscript_id, decision["message_id"], decision["suggestion_id"],
            decision["outcome"], operation_id, _json(decision["request"]), _json(decision["provenance"]), _now(),
        ))

    def decide_suggestion(
        self, manuscript_id: str, message_id: str, *, suggestion_id: str, decision_id: str,
        outcome: str, base_revision: int, selected_change_ids: list[str],
        confirmed_changed_ids: list[str], comment_ids: list[str],
    ) -> dict[str, Any]:
        if outcome not in {"ACCEPT_ALL", "ACCEPT_PARTIAL", "REJECT"}:
            raise ValueError("Unknown author suggestion decision")
        request = {"manuscript_id": manuscript_id, "message_id": message_id, "suggestion_id": suggestion_id,
                   "outcome": outcome, "base_revision": base_revision,
                   "selected_change_ids": sorted(set(selected_change_ids)),
                   "confirmed_changed_ids": sorted(set(confirmed_changed_ids)), "comment_ids": sorted(set(comment_ids))}
        with transaction(self.db_path) as conn:
            previous = conn.execute("SELECT * FROM studio_suggestion_decisions WHERE decision_id=?", (decision_id,)).fetchone()
        if previous is not None:
            if json.loads(previous["request_json"]) != request:
                raise RevisionConflictError("This decision id already belongs to another author choice")
            return {"decision": _decode(previous, "request_json", "provenance_json"),
                    "workspace": self.get_workspace(manuscript_id), "idempotent_replay": True}
        message, suggestion = self._suggestion(manuscript_id, message_id, suggestion_id)
        all_ids = {item["change_id"] for item in suggestion["changes"]}
        selected = set(selected_change_ids)
        if not selected.issubset(all_ids) or not set(confirmed_changed_ids).issubset(selected):
            raise ValueError("Selected changes must belong to this original suggestion")
        if outcome == "ACCEPT_ALL" and selected != all_ids:
            raise ValueError("Full adoption must include every original change")
        if outcome == "ACCEPT_PARTIAL" and (not selected or selected == all_ids):
            raise ValueError("Partial adoption requires a nonempty proper subset of changes")
        if outcome == "REJECT" and (selected or confirmed_changed_ids or comment_ids):
            raise ValueError("Rejection cannot apply changes or associate an edit")
        changes = [item for item in suggestion["changes"] if item["change_id"] in selected]
        with transaction(self.db_path) as conn:
            # Bind the actual author revision, patch and decision in one transaction.
            conn.execute("BEGIN IMMEDIATE")
            previous = conn.execute("SELECT * FROM studio_suggestion_decisions WHERE decision_id=?", (decision_id,)).fetchone()
            replay = previous is not None
            if previous is not None:
                if json.loads(previous["request_json"]) != request:
                    raise RevisionConflictError("This decision id already belongs to another author choice")
            else:
                current = conn.execute("SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,)).fetchone()
                if current is None:
                    raise NotFoundError("Unknown Studio author document")
                document = json.loads(current["document_json"])
                decision = {"decision_id": decision_id, "message_id": message_id, "suggestion_id": suggestion_id,
                            "outcome": outcome, "request": request,
                            "provenance": {"source": message["quoted"]["source"], "suggestion": suggestion,
                                           "model_contract": message["quoted"].get("model_contract"),
                                           "changes": [preview_change(document, item) for item in changes],
                                           "author_revision_at_decision": current["revision"]}}
                if outcome == "REJECT":
                    self._insert_suggestion_decision(conn, manuscript_id, decision, None)
                else:
                    if current["revision"] != base_revision:
                        raise RevisionConflictError("Author revision changed; compare the suggestion again")
                    earlier = conn.execute("SELECT request_json FROM studio_suggestion_decisions WHERE manuscript_id=? AND message_id=? AND suggestion_id=? AND outcome!='REJECT'",
                                           (manuscript_id, message_id, suggestion_id))
                    adopted = {change_id for row in earlier for change_id in json.loads(row[0])["selected_change_ids"]}
                    if selected & adopted:
                        raise ValueError("This original change has already been adopted; use Undo or Redo to reverse an author edit")
                    for item in changes:
                        if item["action"] == "ADD_SOURCE" and not self._suggestion_source_available(conn, manuscript_id, item["evidence"]["evidence_id"]):
                            raise ValueError("An original Evidence candidate is no longer available in this manuscript's scope")
                    patched = apply_suggestion(document, changes, set(confirmed_changed_ids))
                    self.save_document(manuscript_id, base_revision=base_revision, operation_id=decision_id,
                                       operation_kind="ADOPT_STUDIO_SUGGESTION", document=patched,
                                       comment_ids=comment_ids, suggestion_decision=decision, _connection=conn)
            row = conn.execute("SELECT * FROM studio_suggestion_decisions WHERE decision_id=?", (decision_id,)).fetchone()
        return {"decision": _decode(row, "request_json", "provenance_json"),
                "workspace": self.get_workspace(manuscript_id), "idempotent_replay": replay}

    def discuss(
        self,
        manuscript_id: str,
        *,
        content: str,
        anchor_type: str,
        anchor_id: str | None,
        anchor_revision: int,
        quote: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        if not content.strip():
            raise ValueError("Message is empty")
        context = self._chat_anchor(manuscript_id, anchor_type, anchor_id, anchor_revision, quote)
        if "generated" in context and context["generated"]["input_revision"] is not None:
            context["snapshot"] = self.get_document_revision(manuscript_id, context["generated"]["input_revision"])
        quotation = context["quote"]
        user_id = new_id("message")
        timestamp = _now()
        with transaction(self.db_path) as conn:
            conn.execute(
                """INSERT INTO studio_chat_messages VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    user_id,
                    manuscript_id,
                    None,
                    "AUTHOR",
                    content.strip(),
                    anchor_type,
                    anchor_id,
                    anchor_revision,
                    _json(quotation),
                    "COMPLETE",
                    timestamp,
                ),
            )
        sources = self.search_sources(manuscript_id, content, limit=8)
        try:
            response, call_id = self._model_call(
                manuscript_id, "writing_chat",
                {
                    "author_question": content,
                    "anchor": {
                        "type": anchor_type,
                        "id": anchor_id,
                        "revision": anchor_revision,
                        "quote": quotation,
                    },
                    "document": context.get("snapshot", {}).get("document"),
                    "generated_version": context.get("generated"),
                    "retrieved_evidence": sources,
                    "suggestion_format": MODEL_FORMAT,
                    "constraint": "Suggestions must not be applied to the manuscript.",
                }, input_revision=anchor_revision, message_id=user_id,
            )
            agent_content = str(response.get("answer") or response.get("suggestion") or "").strip()
            if not agent_content:
                raise RSTV4Error("Scholar returned no suggestion")
            suggestions = normalize_suggestions(response.get("suggestions", []), context.get("snapshot", {}).get("document"), sources)
        except Exception as exc:
            with transaction(self.db_path) as conn:
                conn.execute(
                    "UPDATE studio_chat_messages SET status='AGENT_FAILED' WHERE message_id=?",
                    (user_id,),
                )
            if isinstance(exc, GatewayDeferredError):
                raise
            raise RSTV4Error(f"Scholar suggestion failed: {exc}") from exc
        agent_id = new_id("message")
        with transaction(self.db_path) as conn:
            conn.execute(
                """INSERT INTO studio_chat_messages VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    agent_id,
                    manuscript_id,
                    user_id,
                    "AGENT",
                    agent_content,
                    anchor_type,
                    anchor_id,
                    anchor_revision,
                    _json({**quotation, "evidence_ids": response.get("evidence_ids", []),
                           "_model_provenance": {"model_contract": response.get("_contract"),
                                                 "model_call_id": call_id},
                           **({"suggestions": suggestions, "model_contract": response.get("_contract")}
                              if suggestions else {})}),
                    "SUGGESTION",
                    _now(),
                ),
            )
            agent_row = conn.execute(
                "SELECT * FROM studio_chat_messages WHERE message_id=?", (agent_id,)
            ).fetchone()
            user_row = conn.execute(
                "SELECT * FROM studio_chat_messages WHERE message_id=?", (user_id,)
            ).fetchone()
        return {"author": self._message(user_row), "agent": self._message(agent_row)}

    def create_job(
        self,
        manuscript_id: str,
        job_kind: str,
        *,
        input_revision: int,
        request: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        if job_kind not in {"GENERATE", "WHOLE_PAPER_REVIEW", "SUBMISSION_PACKAGE"}:
            raise ValueError(f"Unknown Studio job kind: {job_kind}")
        workspace = self.get_workspace(manuscript_id)
        if workspace["document_revision"] != input_revision:
            raise RevisionConflictError("Generation input revision changed")
        if job_kind == "GENERATE":
            self._check_generate_ready(manuscript_id, workspace["document"])
            if workspace["manuscript"]["needs_scientific_review"]:
                raise RSTV4Error("Review the manuscript's scientific dependencies before generation")
        request = json.loads(_json(request or {}))
        if job_kind == "WHOLE_PAPER_REVIEW":
            target_id = request.get("generated_version_id") or workspace["active_generated_version_id"]
            if not target_id:
                raise RSTV4Error("Generate the manuscript before whole-paper review")
            if not isinstance(target_id, str):
                raise ValueError("The review generated version ID must be a string")
            self.get_generated_version(manuscript_id, target_id)
            request["generated_version_id"] = target_id
        job_id = new_id("studio_job")
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            if job_kind == "GENERATE":
                request["scientific_binding"] = scientific_binding_stamp(conn, manuscript_id)
                if request["scientific_binding"]["needs_scientific_review"]:
                    raise RSTV4Error("Review the manuscript's scientific dependencies before generation")
            conn.execute(
                """INSERT INTO studio_jobs VALUES(?,?,?,'QUEUED',?,?,?,?,?,?,?)""",
                (
                    job_id,
                    manuscript_id,
                    job_kind,
                    input_revision,
                    _json(request),
                    None,
                    None,
                    _now(),
                    None,
                    None,
                ),
            )
        self.run_job(job_id)
        return self.get_job(job_id)

    def run_job(self, job_id: str) -> None:
        with self._run_lock:
            if job_id in self._running:
                return
            self._running.add(job_id)
        self._executor.submit(self._execute_job, job_id)

    def retry_queued_jobs(self) -> int:
        with transaction(self.db_path) as conn:
            ids = [row["job_id"] for row in conn.execute(
                "SELECT job_id FROM studio_jobs WHERE status='QUEUED' ORDER BY created_at"
            )]
        for job_id in ids:
            self.run_job(job_id)
        return len(ids)

    def get_job(self, job_id: str) -> dict[str, Any]:
        with transaction(self.db_path) as conn:
            row = conn.execute("SELECT * FROM studio_jobs WHERE job_id=?", (job_id,)).fetchone()
            if row is None:
                raise NotFoundError(f"Unknown Studio job: {job_id}")
        return self._job(row)

    def cancel_job(self, job_id: str) -> dict[str, Any]:
        with transaction(self.db_path) as conn:
            row = conn.execute("SELECT * FROM studio_jobs WHERE job_id=?", (job_id,)).fetchone()
            if row is None:
                raise NotFoundError(f"Unknown Studio job: {job_id}")
            if row["status"] in {"QUEUED", "RUNNING"}:
                conn.execute(
                    "UPDATE studio_jobs SET status='CANCEL_REQUESTED' WHERE job_id=?", (job_id,)
                )
        return self.get_job(job_id)

    def create_review_round(
        self,
        manuscript_id: str,
        *,
        round_number: int,
        venue: str | None,
        baseline_manuscript_version_id: str | None = None,
        expected_document_revision: int | None = None,
    ) -> dict[str, Any]:
        if not isinstance(round_number, int) or isinstance(round_number, bool) or round_number < 1:
            raise ValueError("Review round numbers must be positive integers")
        review_round_id = new_id("review_round")
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            existing = conn.execute("SELECT * FROM review_rounds WHERE manuscript_id=? AND round_number=?", (manuscript_id, round_number)).fetchone()
            if existing:
                result = resolve_review_round(conn, existing)
                if existing["venue"] != venue:
                    raise RevisionConflictError("This round already has a different venue")
                if baseline_manuscript_version_id is not None and result["baseline_manuscript_version_id"] != baseline_manuscript_version_id:
                    raise RevisionConflictError("This round already has a different fixed reviewed version")
                return result
            document = conn.execute("SELECT * FROM studio_documents WHERE manuscript_id=?", (manuscript_id,)).fetchone()
            if document is None:
                raise NotFoundError("Unknown Studio manuscript")
            if expected_document_revision is not None and expected_document_revision != document["revision"]:
                raise RevisionConflictError("Author input changed before the review round was created")
            version_id = baseline_manuscript_version_id or document["active_generated_version_id"]
            original = review_baseline_version(conn, manuscript_id, version_id) if version_id else None
            inputs = json.loads(original["inputs_json"]) if original else {}
            artifact = conn.execute("SELECT artifact_id FROM studio_artifacts WHERE manuscript_id=? AND job_id=? AND artifact_kind='MARKDOWN'", (manuscript_id, inputs.get("job_id"))).fetchone()
            baseline = artifact["artifact_id"] if artifact else None
            conn.execute(
                """INSERT INTO review_rounds VALUES(?,?,?,?,?,'ACTIVE',?)""",
                (review_round_id, manuscript_id, round_number, baseline, venue, _now()),
            )
            conn.execute("INSERT INTO studio_review_round_baselines VALUES(?,?,?,?)", (review_round_id, version_id, content_digest(original["content"]) if original else None, _now()))
            row = conn.execute(
                "SELECT * FROM review_rounds WHERE manuscript_id=? AND round_number=?",
                (manuscript_id, round_number),
            ).fetchone()
            result = resolve_review_round(conn, row)
        return result

    def list_review_rounds(self, manuscript_id: str) -> list[dict[str, Any]]:
        self.get_workspace(manuscript_id)
        with transaction(self.db_path) as conn:
            rows = conn.execute("SELECT * FROM review_rounds WHERE manuscript_id=? ORDER BY round_number", (manuscript_id,)).fetchall()
            return [resolve_review_round(conn, row) for row in rows]

    def record_submission_event(
        self,
        manuscript_id: str,
        *,
        venue: str,
        round_number: int,
        status: str,
        detail: str | None,
    ) -> dict[str, Any]:
        event = self.track_service.record_submission(
            manuscript_id,
            SubmissionEventInput(
                venue=venue,
                round=round_number,
                status=status,
                detail=detail,
            ),
        )
        if status == "PUBLISHED":
            manuscript = self.track_service.get_manuscript(manuscript_id)
            if manuscript.get("workflow_stage") != "PUBLICATION":
                with transaction(self.db_path) as conn:
                    conn.execute(
                        """UPDATE manuscripts SET workflow_stage='PUBLICATION',revision=revision+1,
                           updated_at=? WHERE manuscript_id=?""",
                        (_now(), manuscript_id),
                    )
        return event

    def import_comments(
        self,
        review_round_id: str,
        comments: list[dict[str, Any]],
        *,
        source: str,
        import_metadata: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        selected_count = sum(
            item.get("include") is not False and bool(str(item.get("text") or "").strip())
            for item in comments
        )
        if selected_count > 100:
            raise ValueError("At most 100 selected review comments may be imported in one batch")
        with transaction(self.db_path) as conn:
            round_row = conn.execute(
                "SELECT * FROM review_rounds WHERE review_round_id=?", (review_round_id,)
            ).fetchone()
            if round_row is None:
                raise NotFoundError(f"Unknown review round: {review_round_id}")
            number = self._number_comments(conn, round_row["manuscript_id"])
            timestamp = _now()
            created: list[str] = []
            batch_id = new_id("comment_import")
            for index, item in enumerate(comments):
                text = str(item.get("text") or "").strip()
                if not text or item.get("include") is False:
                    continue
                comment_id = new_id("comment")
                number += 1
                metadata = {
                    **(import_metadata or {}),
                    "raw_import_text": (import_metadata or {}).get(
                        "raw_import_text", (import_metadata or {}).get("original_text")
                    ),
                    "import_batch_id": batch_id,
                    "original_index": index,
                    "original_text": item.get("original_text", text),
                    "comment_number": number,
                }
                conn.execute(
                    """INSERT INTO review_comments VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        comment_id,
                        review_round_id,
                        source,
                        item.get("reviewer"),
                        str(item.get("original_text") or text),
                        text,
                        None,
                        None,
                        "PENDING",
                        None,
                        _json(metadata),
                        timestamp,
                        timestamp,
                    ),
                )
                created.append(comment_id)
            rows = [conn.execute(
                "SELECT * FROM review_comments WHERE comment_id=?", (comment_id,)
            ).fetchone() for comment_id in created]
        return [self._comment(row) for row in rows]

    def list_comments(self, manuscript_id: str, *, include_archived: bool = False) -> list[dict[str, Any]]:
        archive_clause = "" if include_archived else "AND c.archived_at IS NULL"
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            self._number_comments(conn, manuscript_id)
            rows = conn.execute(
                f"""SELECT c.* FROM review_comments c JOIN review_rounds r
                    ON r.review_round_id=c.review_round_id
                    WHERE r.manuscript_id=? {archive_clause}
                    ORDER BY r.round_number,c.created_at,c.rowid""",
                (manuscript_id,),
            ).fetchall()
            results = []
            for row in rows:
                item = self._response_comment(conn, manuscript_id, row)
                round_row = conn.execute("SELECT round_number FROM review_rounds WHERE review_round_id=?", (row["review_round_id"],)).fetchone()
                item["round_number"] = round_row["round_number"]
                item["changes"] = [dict(link) for link in conn.execute(
                    """SELECT sc.* FROM studio_changes sc JOIN comment_changes cc
                       ON cc.change_id=sc.change_id WHERE cc.comment_id=? ORDER BY sc.created_at""",
                    (row["comment_id"],),
                )]
                results.append(item)
        return results

    def _response_comment(self, conn: Any, manuscript_id: str, row: Any) -> dict[str, Any]:
        item = self._comment(row)
        revision = response_checks.response_revision(conn, row["comment_id"])
        item["content_sha256"] = response_checks.content_digest(row, revision)
        item["response_state_sha256"] = response_checks.digest(response_checks.state(conn, manuscript_id, row["comment_id"]))
        item["response_verification"] = response_checks.status(conn, manuscript_id, row)
        return item

    @staticmethod
    def _number_comments(conn: Any, manuscript_id: str) -> int:
        """Backfill stable labels once, including archives, without changing source fields."""
        rows = conn.execute(
            """SELECT c.comment_id,c.import_metadata_json FROM review_comments c
               JOIN review_rounds r ON r.review_round_id=c.review_round_id
               WHERE r.manuscript_id=? ORDER BY r.round_number,c.created_at,c.rowid""",
            (manuscript_id,),
        ).fetchall()
        metadata = [(row, json.loads(row["import_metadata_json"])) for row in rows]
        number = max((item["comment_number"] for _, item in metadata
                      if isinstance(item.get("comment_number"), int)
                      and item["comment_number"] > 0), default=0)
        for row, item in metadata:
            if not isinstance(item.get("comment_number"), int) or item["comment_number"] <= 0:
                number += 1
                conn.execute(
                    "UPDATE review_comments SET import_metadata_json=? WHERE comment_id=?",
                    (_json({**item, "comment_number": number}), row["comment_id"]),
                )
        return number

    def update_comment(self, comment_id: str, changes: dict[str, Any]) -> dict[str, Any]:
        allowed = {"edited_text", "author_decision", "response_text", "status"}
        fields = {key: value for key, value in changes.items() if key in allowed}
        if not fields:
            raise ValueError("No supported review comment fields supplied")
        if "status" in fields and fields["status"] not in {
            "PENDING",
            "DECIDED",
            "MODIFYING",
            "VERIFY",
            "RESOLVED",
        }:
            raise ValueError("Unknown review comment status")
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            self._number_comment_owner(conn, comment_id)
            original = conn.execute(
                """SELECT c.*,r.manuscript_id FROM review_comments c JOIN review_rounds r
                   ON r.review_round_id=c.review_round_id WHERE c.comment_id=?""", (comment_id,),
            ).fetchone()
            if original is None:
                raise NotFoundError(f"Unknown review comment: {comment_id}")
            current = self._response_comment(conn, original["manuscript_id"], original)
            if changes.get("expected_content_sha256") not in {None, current["content_sha256"]}:
                raise RevisionConflictError("This opinion or reply changed elsewhere; keep your draft and compare the current reply")
            changed = any(original[key] != value for key, value in fields.items())
            if fields.get("status") == "RESOLVED" and (changed or current["response_verification"]["status"] != "CURRENT"):
                raise ValueError("Use the explicit author reply review to mark this opinion complete")
            if changed:
                response_checks.advance(conn, comment_id)
                if original["status"] == "RESOLVED" and "status" not in fields:
                    fields["status"] = "VERIFY"
            assignments = ",".join(f"{key}=?" for key in fields)
            conn.execute(
                f"UPDATE review_comments SET {assignments},updated_at=? WHERE comment_id=?",
                (*fields.values(), _now(), comment_id),
            )
            row = conn.execute(
                "SELECT * FROM review_comments WHERE comment_id=?", (comment_id,)
            ).fetchone()
            result = self._response_comment(conn, original["manuscript_id"], row)
        return result

    def response_review_context(self, manuscript_id: str, comment_id: str) -> dict[str, Any]:
        """Read exact saved reply, immutable generation and actual compiled layout."""
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN")
            current = response_checks.state(conn, manuscript_id, comment_id)
            row = conn.execute("SELECT * FROM review_comments WHERE comment_id=?", (comment_id,)).fetchone()
            comment = self._comment(row)
            comment["changes"] = [dict(item) for item in conn.execute(
                """SELECT sc.* FROM studio_changes sc JOIN comment_changes cc ON sc.change_id=cc.change_id
                   WHERE cc.comment_id=? ORDER BY sc.change_id""", (comment_id,),
            )]
        if not current["generated_version_id"]:
            raise ValueError("Generate the current manuscript before reviewing its final reply")
        target = self.get_generated_version(manuscript_id, current["generated_version_id"])
        baseline = (self.get_generated_version(manuscript_id, current["baseline_manuscript_version_id"])
                    if current["baseline_manuscript_version_id"] else None)
        target_document = (self.get_document_revision(manuscript_id, target["input_revision"])["document"]
                           if target["input_revision"] is not None else None)
        baseline_document = (self.get_document_revision(manuscript_id, baseline["input_revision"])["document"]
                             if baseline and baseline["input_revision"] is not None else None)
        pdfs = [item for item in target["artifacts"] if item["artifact_kind"] == "PDF"]
        evidence, layout = [], {"status": "PDF_UNAVAILABLE", "lines": []}
        if len(pdfs) == 1 and target_document is not None:
            raw = self._verified_artifact_bytes(pdfs[0])
            with tempfile.TemporaryDirectory(prefix="studio-reply-review-") as directory:
                folder = Path(directory)
                pdf = folder / "fixed-manuscript.pdf"
                pdf.write_bytes(raw)
                write_mapping(folder / "change-location-map.csv", [comment], target=target,
                              baseline=baseline, target_document=target_document, baseline_document=baseline_document, pdf=pdf)
                mapping = json.loads((folder / "change-realization-map.json").read_text(encoding="utf-8"))
                evidence = [{**item, "evidence_id": response_checks.digest(item)} for item in mapping["rows"]]
                layout = mapping["pdf"]
        before = baseline["content"] if baseline else ""
        context = {
            "state": current, "state_sha256": response_checks.digest(current),
            "baseline_available": baseline is not None,
            "baseline_version_number": baseline["version"] if baseline else None,
            "target_version_number": target["version"],
            "baseline_text_sha256": content_digest(before) if baseline else None,
            "target_text_sha256": content_digest(target["content"]),
            "baseline_to_target_diff": "\n".join(difflib.unified_diff(before.splitlines(), target["content"].splitlines(),
                                         fromfile="reviewed-baseline.md", tofile="final-manuscript.md", lineterm="")) if baseline else "",
            "evidence": evidence, "pdf": layout,
            "pdf_artifact_id": pdfs[0]["artifact_id"] if len(pdfs) == 1 else None,
            "scientific_acceptance": False, "automatic_response_approval": False,
        }
        context["context_sha256"] = response_checks.digest(context)
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN")
            if response_checks.digest(response_checks.state(conn, manuscript_id, comment_id)) != context["state_sha256"]:
                raise RevisionConflictError("Reply context changed while loading; reopen the review")
            context["history"] = response_checks.history(conn, comment_id)
            context["verification"] = response_checks.status(conn, manuscript_id, row)
        return context

    def confirm_response_review(self, manuscript_id: str, comment_id: str, request: dict[str, Any]) -> dict[str, Any]:
        # A retry returns the original decision even when later changes made it stale.
        with transaction(self.db_path) as conn:
            response_checks.state(conn, manuscript_id, comment_id)
            previous = conn.execute("SELECT * FROM studio_response_checks WHERE request_id=?", (request["request_id"],)).fetchone()
            frozen = json.loads(previous["context_json"]) if previous else None
        context = frozen if frozen is not None else self.response_review_context(manuscript_id, comment_id)
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            result = response_checks.record(conn, manuscript_id, comment_id, context, request)
            row = conn.execute("SELECT * FROM review_comments WHERE comment_id=?", (comment_id,)).fetchone()
            return {**result, "comment": self._response_comment(conn, manuscript_id, row),
                    "history": response_checks.history(conn, comment_id)}

    def response_review_history(self, manuscript_id: str, comment_id: str) -> dict[str, Any]:
        """Historical replies remain readable without generating or reopening a PDF."""
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN")
            row = conn.execute("""SELECT c.* FROM review_comments c JOIN review_rounds r
                                  ON r.review_round_id=c.review_round_id WHERE c.comment_id=? AND r.manuscript_id=?""",
                               (comment_id, manuscript_id)).fetchone()
            if row is None:
                raise NotFoundError("The review comment does not belong to this manuscript")
            return {"history": response_checks.history(conn, comment_id)}

    def archive_comment(self, comment_id: str, *, restore: bool = False) -> dict[str, Any]:
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            self._number_comment_owner(conn, comment_id)
            if conn.execute(
                "SELECT 1 FROM review_comments WHERE comment_id=?", (comment_id,)
            ).fetchone() is None:
                raise NotFoundError(f"Unknown review comment: {comment_id}")
            conn.execute(
                "UPDATE review_comments SET archived_at=?,updated_at=? WHERE comment_id=?",
                (None if restore else _now(), _now(), comment_id),
            )
            response_checks.advance(conn, comment_id)
            row = conn.execute(
                "SELECT * FROM review_comments WHERE comment_id=?", (comment_id,)
            ).fetchone()
        return self._comment(row)

    def _number_comment_owner(self, conn: Any, comment_id: str) -> None:
        owner = conn.execute(
            """SELECT r.manuscript_id FROM review_comments c JOIN review_rounds r
               ON r.review_round_id=c.review_round_id WHERE c.comment_id=?""", (comment_id,),
        ).fetchone()
        if owner is not None:
            self._number_comments(conn, owner["manuscript_id"])

    def _ensure_document(self, manuscript: dict[str, Any], *, seed_new: bool = False) -> Any:
        with transaction(self.db_path) as conn:
            row = conn.execute(
                "SELECT * FROM studio_documents WHERE manuscript_id=?",
                (manuscript["manuscript_id"],),
            ).fetchone()
            if row is None:
                timestamp = _now()
                document = _default_document(manuscript) if seed_new else _legacy_document(manuscript)
                latest_version_id = (
                    manuscript["versions"][-1]["manuscript_version_id"]
                    if manuscript.get("versions") and not seed_new
                    else None
                )
                conn.execute(
                    "INSERT INTO studio_documents VALUES(?,0,?,?,NULL,?,?)",
                    (
                        manuscript["manuscript_id"],
                        _json(document),
                        latest_version_id,
                        timestamp,
                        timestamp,
                    ),
                )
                row = conn.execute(
                    "SELECT * FROM studio_documents WHERE manuscript_id=?",
                    (manuscript["manuscript_id"],),
                ).fetchone()
        return row

    @staticmethod
    def _validate_document_shape(manuscript_id: str, document: dict[str, Any]) -> None:
        if document.get("schema_version") != 1:
            raise ValueError("Unsupported Studio document schema")
        if document.get("manuscript_id") != manuscript_id:
            raise ValueError("Document manuscript id does not match route")
        if not isinstance(document.get("sections"), list):
            raise TypeError("Document sections must be a list")
        seen: set[str] = {manuscript_id}
        for section in document["sections"]:
            if section.get("placement_status") not in {None, "UNPLACED"}:
                raise ValueError("Unknown section placement status")
            for item in [section, *section.get("paragraphs", [])]:
                identifier = item.get("id")
                if not identifier or identifier in seen:
                    raise ValueError("Every structural block must have a unique stable id")
                seen.add(identifier)
            for paragraph in section.get("paragraphs", []):
                if not isinstance(paragraph.get("logical_relation", ""), str):
                    raise TypeError("Paragraph logical relation must be author text")
                for formula in paragraph.get("formulas", []):
                    if not formula.get("id") or formula["id"] in seen:
                        raise ValueError("Every formula must have a unique stable id")
                    seen.add(formula["id"])
                    latex = str(formula.get("latex") or "")
                    if latex:
                        StudioService._validate_formula(latex)
                for intent in paragraph.get("intents", []):
                    identifier = intent.get("id")
                    if not identifier or identifier in seen:
                        raise ValueError("Every intent must have a unique stable id")
                    seen.add(identifier)
                    for thought in intent.get("thoughts", []):
                        identifier = thought.get("id")
                        if not identifier or identifier in seen:
                            raise ValueError("Every thought must have a unique stable id")
                        seen.add(identifier)
                        mode = thought.get("citation_mode", "PENDING")
                        if mode not in {"PENDING", "NONE", "SELECTED"}:
                            raise ValueError("Unknown thought citation mode")
        validate_groups(document, seen)
        section_levels(document["sections"])
        validate_argument_relations(document, seen)
        validate_correspondence(document)

    @staticmethod
    def _changed_block_ids(before: dict[str, Any], after: dict[str, Any]) -> set[str]:
        def blocks(document: dict[str, Any]) -> dict[str, str]:
            author_owners, context_owners = validate_correspondence(document)
            values: dict[str, str] = {"paper": _json({
                "title": document.get("title"),
                "abstract": document.get("abstract"),
                "venue": document.get("venue"),
            })}
            for section_index, section in enumerate(document.get("sections", [])):
                values[section["id"]] = _json({
                    "title": section.get("title"),
                    "role": section.get("role"),
                    "placement_status": section.get("placement_status"),
                    "parent_section_id": section.get("parent_section_id"),
                    "block_order": section.get("block_order"),
                    "paragraph_ids": [item["id"] for item in section.get("paragraphs", [])],
                    "position": section_index,
                })
                for group in section.get("paragraph_groups", []):
                    values[group["id"]] = _json({**group, "section_id": section["id"]})
                for paragraph in section.get("paragraphs", []):
                    values[paragraph["id"]] = _json({
                        "purpose": paragraph.get("purpose"),
                        "paradigm": paragraph.get("paradigm"),
                        "logic": paragraph.get("logic"),
                        "logical_relation": paragraph.get("logical_relation"),
                        "template_source": paragraph.get("template_source"),
                        "author_context_paragraph_ids": paragraph.get("author_context_paragraph_ids"),
                        "context_target_paragraph_id": context_owners.get(paragraph["id"]),
                        "section_id": section["id"],
                        "group_id": (paragraph_group(section, paragraph["id"]) or {}).get("id"),
                        "intent_ids": [item["id"] for item in paragraph.get("intents", [])],
                    })
                    for intent in paragraph.get("intents", []):
                        values[intent["id"]] = _json({"label": intent.get("label"), "paragraph_id": paragraph["id"], "thought_ids": [item["id"] for item in intent.get("thoughts", [])],
                            "author_intent_ids": intent.get("author_intent_ids"),
                            "template_target_intent_id": author_owners.get(intent["id"])})
                        for thought in intent.get("thoughts", []):
                            values[thought["id"]] = _json({**thought, "intent_id": intent["id"]})
            for relation in document.get("argument_relations", []):
                values[relation["id"]] = _json(relation)
            return values

        old = blocks(before)
        new = blocks(after)
        return {
            identifier
            for identifier in old.keys() | new.keys()
            if old.get(identifier) != new.get(identifier)
        }

    def _check_generate_ready(self, manuscript_id: str, document: dict[str, Any]) -> None:
        issues: list[dict[str, str]] = []
        if has_correspondence(document):
            issues.extend(unresolved_correspondence(document))
            document = realization_document(document)
        if not str(document.get("title") or "").strip():
            issues.append({"path": "title", "message": "Paper title is required"})
        thought_count = 0
        asset_ids: set[str] = set()
        for section in document.get("sections", []):
            if section.get("placement_status") == "UNPLACED":
                issues.append({"path": section["id"],
                               "message": "Original author content remains unplaced"})
            if not str(section.get("title") or "").strip():
                issues.append({"path": section.get("id", "section"), "message": "Section title is missing"})
            if not section.get("paragraphs") and not any(
                item.get("parent_section_id") == section["id"] for item in document["sections"]
            ):
                issues.append({"path": section["id"], "message": "Section has no paragraph"})
            for group in section.get("paragraph_groups", []):
                if not group["paragraph_ids"] or not group["label"].strip():
                    issues.append({"path": group["id"], "message": "Paragraph group is empty or its author label is unresolved"})
            for paragraph in section.get("paragraphs", []):
                if not paragraph.get("intents"):
                    issues.append({"path": paragraph["id"], "message": "Paragraph intent is unresolved"})
                asset_ids.update(paragraph.get("asset_ids", []))
                for formula in paragraph.get("formulas", []):
                    latex = str(formula.get("latex") or "").strip()
                    if not latex:
                        issues.append({"path": formula["id"], "message": "Formula LaTeX is empty"})
                    else:
                        try:
                            self._validate_formula(latex)
                        except ValueError as exc:
                            issues.append({"path": formula["id"], "message": str(exc)})
                for intent in paragraph.get("intents", []):
                    if not str(intent.get("label") or "").strip():
                        issues.append({"path": intent["id"], "message": "Intent label is unresolved"})
                    if not intent.get("thoughts"):
                        issues.append({"path": intent["id"], "message": "Intent has no author thought"})
                    for thought in intent.get("thoughts", []):
                        thought_count += 1
                        if not str(thought.get("text") or "").strip():
                            issues.append({"path": thought["id"], "message": "Thought is empty"})
                        mode = thought.get("citation_mode", "PENDING")
                        selected = [item for item in thought.get("citations", []) if item.get("selected")]
                        if mode == "PENDING":
                            issues.append({"path": thought["id"], "message": "Citation decision is pending"})
                        if mode == "SELECTED" and not selected:
                            issues.append({"path": thought["id"], "message": "No citation is selected"})
        if thought_count == 0:
            issues.append({"path": "sections", "message": "At least one author thought is required"})
        if asset_ids:
            with transaction(self.db_path) as conn:
                placeholders = ",".join("?" for _ in asset_ids)
                ready = {row["asset_id"] for row in conn.execute(
                    f"SELECT asset_id FROM studio_assets WHERE manuscript_id=? AND status='READY' "
                    f"AND asset_id IN ({placeholders})",
                    (manuscript_id, *asset_ids),
                )}
            for asset_id in sorted(asset_ids - ready):
                issues.append({"path": asset_id, "message": "Attachment is unavailable or not ready"})
            from copy import deepcopy
            for asset_id in sorted(asset_ids & ready):
                probe = deepcopy(document)
                probe["sections"] = [{"paragraphs": [{"asset_ids": [asset_id]}]}]
                try:
                    with transaction(self.db_path) as conn:
                        freeze_materials(conn, self.storage_root, manuscript_id, probe)
                except (ValueError, OSError, ImportError) as exc:
                    issues.append({"path": asset_id, "message": f"Original material cannot be rendered: {exc}"})
        if issues:
            raise DocumentValidationError(issues)

    def _execute_job(self, job_id: str) -> None:
        try:
            with transaction(self.db_path) as conn:
                row = conn.execute(
                    "SELECT * FROM studio_jobs WHERE job_id=?", (job_id,)
                ).fetchone()
                if row is None or row["status"] != "QUEUED":
                    return
                conn.execute(
                    "UPDATE studio_jobs SET status='RUNNING',started_at=? WHERE job_id=?",
                    (_now(), job_id),
                )
                job = dict(row)
            if job["job_kind"] == "GENERATE":
                result = self._generate(job)
            elif job["job_kind"] == "WHOLE_PAPER_REVIEW":
                result = self._review(job)
            else:
                result = self._submission_package(job)
            with transaction(self.db_path) as conn:
                current = conn.execute(
                    "SELECT status FROM studio_jobs WHERE job_id=?", (job_id,)
                ).fetchone()
                status = "CANCELLED" if current["status"] == "CANCEL_REQUESTED" else "SUCCEEDED"
                conn.execute(
                    """UPDATE studio_jobs SET status=?,result_json=?,completed_at=?
                       WHERE job_id=?""",
                    (status, _json(result), _now(), job_id),
                )
        except Exception as exc:  # noqa: BLE001 - durable job boundary records every failure
            if isinstance(exc, JobCancelled):
                with transaction(self.db_path) as conn:
                    conn.execute(
                        """UPDATE studio_jobs SET status='CANCELLED',error_json=NULL,
                           completed_at=? WHERE job_id=?""",
                        (_now(), job_id),
                    )
                return
            payload = {"type": type(exc).__name__, "message": str(exc)}
            if isinstance(exc, DocumentValidationError):
                payload["issues"] = exc.issues
            with transaction(self.db_path) as conn:
                conn.execute(
                    """UPDATE studio_jobs SET status='FAILED',error_json=?,completed_at=?
                       WHERE job_id=?""",
                    (_json(payload), _now(), job_id),
                )
        finally:
            with self._run_lock:
                self._running.discard(job_id)

    def _generate(self, job: dict[str, Any]) -> dict[str, Any]:
        workspace = self.get_workspace(job["manuscript_id"])
        with transaction(self.db_path) as conn:
            scientific_binding = scientific_binding_stamp(conn, job["manuscript_id"])
        request = job.get("request") or json.loads(job.get("request_json") or "{}")
        if scientific_binding["needs_scientific_review"]:
            raise RSTV4Error("Review the manuscript's scientific dependencies before generation")
        if request.get("scientific_binding", scientific_binding) != scientific_binding:
            raise RevisionConflictError("Scientific bindings changed after generation was requested")
        if workspace["document_revision"] != job["input_revision"]:
            raise RevisionConflictError("Generation input became stale before execution")
        author_document = workspace["document"]
        document = realization_document(author_document)
        current_stage = workspace["workflow_stage"]
        self._check_generate_ready(job["manuscript_id"], author_document)
        with transaction(self.db_path) as conn:
            frozen_materials, material_files = freeze_materials(
                conn, self.storage_root, job["manuscript_id"], document
            )
        if material_files:
            material_directory = self.storage_root / job["manuscript_id"] / "generated" / job["job_id"]
            material_directory.mkdir(parents=True, exist_ok=True)
            for filename, content in material_files.items():
                (material_directory / filename).write_bytes(content)
        selected_ids = list(dict.fromkeys(
            item["evidence_id"]
            for section in document["sections"] for paragraph in section["paragraphs"]
            for intent in paragraph["intents"] for thought in intent["thoughts"]
            for item in thought.get("citations", [])
            if item.get("selected") and item.get("evidence_id")
        ))
        # One source snapshot feeds realization, references and compilation.
        # A later metadata refresh must not rewrite this generated version.
        frozen_evidence = {item["evidence_id"]: item for item in self._evidence_by_ids(selected_ids)}
        if set(frozen_evidence) != set(selected_ids):
            raise RSTV4Error("A selected source is unavailable")
        track = self.track_service.get_track(workspace["manuscript"]["track_id"])
        if any(
            item["track_id"] != track["track_id"] and item["field_id"] != track["field_id"]
            for item in frozen_evidence.values()
        ):
            raise RSTV4Error("A selected source is outside this manuscript's Field and Track")
        generated_sections: list[dict[str, Any]] = []
        for section in document["sections"]:
            generated_paragraphs = []
            for paragraph in section["paragraphs"]:
                if self._job_cancel_requested(job["job_id"]):
                    raise JobCancelled("Generation cancelled by author")
                evidence_ids = []
                for intent in paragraph["intents"]:
                    for thought in intent["thoughts"]:
                        evidence_ids.extend(
                            item["evidence_id"]
                            for item in thought.get("citations", [])
                            if item.get("selected") and item.get("evidence_id")
                        )
                evidence = [frozen_evidence[item] for item in dict.fromkeys(evidence_ids)]
                payload = {
                        "section_title": section["title"],
                        "section_role": section.get("role"),
                        **({"section_path": section_titles(document, section["id"])}
                           if section.get("parent_section_id") else {}),
                        "paragraph_purpose": paragraph.get("purpose"),
                        "paradigm": paragraph.get("paradigm"),
                        **({"logical_relation": paragraph["logical_relation"]}
                           if paragraph.get("logical_relation") else {}),
                        **({"argument_relations": argument_context(author_document, paragraph)}
                           if argument_context(author_document, paragraph) else {}),
                        **({"author_contexts": paragraph["author_contexts"]}
                           if paragraph.get("author_contexts") else {}),
                        **({"paragraph_group": paragraph_group(section, paragraph["id"])}
                           if paragraph_group(section, paragraph["id"]) else {}),
                        "intents": paragraph["intents"],
                        "selected_evidence": [model_source(item) for item in evidence],
                        **({"assets": [frozen_materials[a] for a in paragraph["asset_ids"]]}
                           if paragraph.get("asset_ids") else {}),
                        "output_schema": OUTPUT_SCHEMA,
                        "constraint": (
                            "Preserve every author thought; do not add scientific claims. "
                            "Use only selected evidence ids in citations."
                        ),
                }
                expected_thought_ids = {
                    thought["id"]
                    for intent in paragraph["intents"]
                    for thought in intent["thoughts"]
                }
                feedback = None
                for attempt in range(3):
                    if self._job_cancel_requested(job["job_id"]):
                        raise JobCancelled("Generation cancelled by author")
                    try:
                        response, call_id = self._model_call(
                            job["manuscript_id"], "paragraph_realization",
                            {**payload, **({"schema_correction": feedback} if feedback else {})},
                            job_id=job["job_id"], input_revision=job["input_revision"],
                            validate=partial(validate_realization, thought_ids=expected_thought_ids,
                                             evidence_ids=set(evidence_ids)),
                        )
                        realization = validate_realization(response, expected_thought_ids, set(evidence_ids))
                        break
                    except RSTV4Error as exc:
                        if attempt == 2:
                            raise
                        feedback = {"attempt": attempt + 1, "error": str(exc),
                                    "requirement": "Return the declared schema without changing author thoughts or adding sources."}
                generated_paragraphs.append(
                    {
                        "paragraph_id": paragraph["id"],
                        "text": realization["text"],
                        "thought_ids": [
                            thought["id"]
                            for intent in paragraph["intents"]
                            for thought in intent["thoughts"]
                        ],
                        "sentence_mappings": realization["sentence_mappings"],
                        "sentences": realization["sentences"],
                        "sentence_positions_verified": realization["sentence_positions_verified"],
                        "model_contract": response.get("_contract"),
                        "model_call_id": call_id,
                        "evidence_ids": sorted(set(evidence_ids)),
                        "formulas": paragraph.get("formulas", []),
                        **({"asset_ids": paragraph["asset_ids"]}
                           if paragraph.get("asset_ids") or paragraph.get("author_contexts") else {}),
                        **({"assets": [frozen_materials[a] for a in paragraph["asset_ids"]]}
                           if paragraph.get("asset_ids") else {}),
                        **({"author_contexts": paragraph["author_contexts"]}
                           if paragraph.get("author_contexts") else {}),
                        **({"group_id": paragraph_group(section, paragraph["id"])["id"]}
                           if paragraph_group(section, paragraph["id"]) else {}),
                    }
                )
            generated_sections.append(
                {"section_id": section["id"], "title": section["title"], "paragraphs": generated_paragraphs,
                 **({"parent_section_id": section["parent_section_id"],
                     "section_depth": section_levels(document["sections"])[section["id"]]}
                    if section.get("parent_section_id") else {}),
                 **({"paragraph_groups": section["paragraph_groups"], "block_order": section["block_order"]}
                    if section.get("paragraph_groups") else {})}
            )
        generated = {
            "title": document["title"],
            "abstract": document.get("abstract", ""),
            "sections": generated_sections,
            "input_revision": job["input_revision"],
            "citation_style": "numeric_first_use",
            **({"material_snapshot_schema": 1} if frozen_materials else {}),
            **({"author_source_sections": document["author_source_sections"]}
               if document.get("author_source_sections") else {}),
            **({"argument_relations": document["argument_relations"]}
               if document.get("argument_relations") else {}),
        }
        generated["references"] = self._reference_snapshot(generated, list(frozen_evidence.values()))
        generated["reference_snapshot_schema"] = 1
        if self._job_cancel_requested(job["job_id"]):
            raise JobCancelled("Generation cancelled by author")
        artifacts = self._compile_artifacts(job, generated)
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            status = conn.execute(
                "SELECT status FROM studio_jobs WHERE job_id=?", (job["job_id"],)
            ).fetchone()["status"]
            if status == "CANCEL_REQUESTED":
                raise JobCancelled("Generation cancelled by author")
            studio = conn.execute(
                "SELECT revision FROM studio_documents WHERE manuscript_id=?",
                (job["manuscript_id"],),
            ).fetchone()
            if studio["revision"] != job["input_revision"]:
                raise RevisionConflictError("Generated output is stale and was not promoted")
            if scientific_binding_stamp(conn, job["manuscript_id"]) != scientific_binding:
                raise RevisionConflictError("Scientific bindings changed during generation; old output was not promoted")
            version = conn.execute(
                "SELECT COALESCE(MAX(version),0)+1 FROM manuscript_versions WHERE manuscript_id=?",
                (job["manuscript_id"],),
            ).fetchone()[0]
            version_id = new_id("manuscript_version")
            markdown = self._generated_markdown(generated)
            conn.execute(
                "INSERT INTO manuscript_versions VALUES(?,?,?,?,?,?)",
                (
                    version_id,
                    job["manuscript_id"],
                    version,
                    markdown,
                    _json({"studio_revision": job["input_revision"], "job_id": job["job_id"],
                           "generated": generated, "scientific_binding": scientific_binding}),
                    _now(),
                ),
            )
            conn.execute(
                """UPDATE studio_documents SET active_generated_version_id=?,
                   last_generated_revision=?,updated_at=? WHERE manuscript_id=?""",
                (version_id, job["input_revision"], _now(), job["manuscript_id"]),
            )
            target_stage = (
                current_stage
                if current_stage in {
                    "WHOLE_PAPER_REVIEW",
                    "FINALIZATION_SUBMISSION",
                    "REVISION",
                    "PUBLICATION",
                }
                else "WHOLE_PAPER_REVIEW"
            )
            conn.execute(
                """UPDATE manuscripts SET workflow_stage=?,writing_status=?,
                   revision=revision+1,updated_at=? WHERE manuscript_id=?""",
                (
                    target_stage,
                    _STATUS_BY_STAGE[target_stage],
                    _now(),
                    job["manuscript_id"],
                ),
            )
        return {"generated_version_id": version_id, "artifacts": artifacts}

    def _review(self, job: dict[str, Any]) -> dict[str, Any]:
        workspace = self.get_workspace(job["manuscript_id"])
        request = job.get("request") or json.loads(job.get("request_json") or "{}")
        version_id = request.get("generated_version_id") or workspace["active_generated_version_id"]
        if not version_id:
            raise RSTV4Error("Generate the manuscript before whole-paper review")
        version = self.get_generated_version(job["manuscript_id"], version_id)
        text = version["content"]
        findings: list[dict[str, Any]] = []
        paragraphs = [item.strip() for item in text.split("\n\n") if len(item.strip()) > 40]
        normalized: dict[str, int] = {}
        for index, paragraph in enumerate(paragraphs):
            key = re.sub(r"\W+", " ", paragraph.casefold()).strip()
            if key in normalized:
                findings.append(
                    {
                        "kind": "DUPLICATE",
                        "severity": "WARNING",
                        "anchor": {"paragraph_index": index, "duplicate_of": normalized[key]},
                        "detail": "The same paragraph text appears more than once.",
                        "source": "RULE",
                    }
                )
            normalized[key] = index
        citation_check = self.track_service.verify_citations(job["manuscript_id"], text)
        for key in citation_check.get("invalid_citations", []):
            findings.append(
                {
                    "kind": "UNRESOLVED_CITATION",
                    "severity": "ERROR",
                    "anchor": {"citation": key},
                    "detail": f"Citation {key} cannot be resolved in the manuscript evidence set.",
                    "source": "RULE",
                }
            )
        semantic_status = "AVAILABLE"
        call_id, model_contract = None, None
        try:
            if version["input_revision"] is None:
                raise RSTV4Error("The historical generated version has no recorded author input revision")
            fixed_document = self.get_document_revision(job["manuscript_id"], version["input_revision"])["document"]
            response, call_id = self._model_call(
                job["manuscript_id"],
                "writing_review",
                {
                    "generated_version_id": version_id,
                    "author_input_revision": version["input_revision"],
                    "generated_manuscript": text,
                    "author_structure": fixed_document,
                    "fixed_references": version["references"],
                    "constraint": "Report possible drift or contradiction; never rewrite the manuscript.",
                    "output_schema": {"findings": [{"kind": "nonempty string", "detail": "nonempty string",
                                                    "severity": "INFO | WARNING | ERROR", "anchor": {}}]},
                },
                input_revision=version["input_revision"], job_id=job["job_id"], validate=_validate_review,
            )
            model_contract = response.get("_contract")
            for item in response.get("findings", []):
                findings.append(
                    {
                        "kind": item.get("kind", "SEMANTIC_CHECK"),
                        "severity": item.get("severity", "WARNING"),
                        "anchor": item.get("anchor", {}),
                        "detail": item.get("detail", "Semantic review finding"),
                        "source": "MODEL",
                    }
                )
        except Exception as exc:  # noqa: BLE001 - model review degrades to explicit rule-only status
            semantic_status = f"UNAVAILABLE: {exc}"
            recorded_calls = self.list_model_calls(job["manuscript_id"], job_id=job["job_id"])
            if recorded_calls:
                call_id = recorded_calls[-1]["call_id"]
                model_contract = recorded_calls[-1]["contract"]
        with transaction(self.db_path) as conn:
            if conn.execute("SELECT status FROM studio_jobs WHERE job_id=?", (job["job_id"],)).fetchone()["status"] == "CANCEL_REQUESTED":
                return {"generated_version_id": version_id, "findings": [], "semantic_status": "CANCELLED",
                        "model_call_id": call_id, "model_contract": model_contract}
            for item in findings:
                conn.execute(
                    """INSERT INTO studio_findings VALUES(?,?,?,?,?,?,?,?,?,?,NULL)""",
                    (
                        new_id("finding"),
                        job["manuscript_id"],
                        version_id,
                        item["kind"],
                        item["severity"],
                        _json(item["anchor"]),
                        item["detail"],
                        item["source"],
                        "OPEN",
                        _now(),
                    ),
                )
        return {"generated_version_id": version_id, "findings": findings, "semantic_status": semantic_status,
                "model_call_id": call_id, "model_contract": model_contract,
                "author_input_revision": version["input_revision"],
                "target_origin": "FIXED_AT_JOB_CREATION" if request.get("generated_version_id")
                else "LEGACY_CURRENT_VERSION_AT_EXECUTION"}

    def _submission_package(self, job: dict[str, Any]) -> dict[str, Any]:
        workspace = self.get_workspace(job["manuscript_id"])
        if workspace["document_revision"] != job["input_revision"]:
            raise RevisionConflictError("Submission package input became stale")
        if workspace["generated_is_stale"]:
            raise RSTV4Error("Regenerate the manuscript before creating a submission package")
        if not workspace["active_generated_version_id"]:
            raise RSTV4Error("Generate the manuscript before creating a submission package")
        request = job.get("request") or json.loads(job.get("request_json") or "{}")
        with transaction(self.db_path) as conn:
            requested_round = request.get("review_round_id")
            review_round = conn.execute(
                "SELECT * FROM review_rounds WHERE review_round_id=? AND manuscript_id=?",
                (requested_round, job["manuscript_id"]),
            ).fetchone() if requested_round else conn.execute(
                """SELECT * FROM review_rounds WHERE manuscript_id=?
                   ORDER BY round_number DESC LIMIT 1""",
                (job["manuscript_id"],),
            ).fetchone()
            if requested_round and review_round is None:
                raise NotFoundError("The selected review round does not belong to this manuscript")
            baseline = ""
            if review_round is not None:
                review_round, baseline = baseline_text(conn, review_round)
        generated_version = self.get_generated_version(job["manuscript_id"], workspace["active_generated_version_id"])
        comments = [item for item in self.list_comments(job["manuscript_id"], include_archived=False)
                    if review_round is not None and item["review_round_id"] == review_round["review_round_id"]]
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN")
            response_snapshot = self._package_response_snapshot(conn, job["manuscript_id"], review_round)
            for comment in comments:
                state = response_snapshot.get(comment["comment_id"])
                if state is None or state["content_sha256"] != comment["content_sha256"] or state["state_sha256"] != comment["response_state_sha256"]:
                    raise RevisionConflictError("Review replies changed while packaging; retry with the current replies")
            if set(response_snapshot) != {item["comment_id"] for item in comments}:
                raise RevisionConflictError("The review opinion list changed while packaging")
            response_history = [{"comment_id": item["comment_id"],
                                 "verification": response_snapshot[item["comment_id"]]["verification"],
                                 "checks": response_checks.history(conn, item["comment_id"])} for item in comments]
            # Use the same snapshot as the exported verification and history.
            for item in comments:
                item["status"] = response_snapshot[item["comment_id"]]["status"]
                item["response_verification"] = response_snapshot[item["comment_id"]]["verification"]
        package_dir = self.storage_root / job["manuscript_id"] / "packages" / job["job_id"]
        package_dir.mkdir(parents=True, exist_ok=True)
        response_lines = ["# Response to reviewers", ""]
        mapping_rows = [["comment_id", "status", "decision", "response", "change_ids"]]
        unresolved = []
        for comment in comments:
            response_lines.extend(
                [
                    f"## {comment['comment_id']}",
                    comment["edited_text"],
                    "",
                    comment.get("response_text") or "Response pending.",
                    "",
                ]
            )
            change_ids = [change["change_id"] for change in comment.get("changes", [])]
            mapping_rows.append(
                [
                    comment["comment_id"],
                    comment["status"],
                    comment.get("author_decision") or "",
                    comment.get("response_text") or "",
                    ";".join(change_ids),
                ]
            )
            if comment["status"] != "RESOLVED" or comment["response_verification"]["status"] != "CURRENT":
                unresolved.append(comment["comment_id"])
        response_path = package_dir / "response-to-reviewers.md"
        response_path.write_text("\n".join(response_lines), encoding="utf-8")
        mapping_path = package_dir / "comment-change-map.csv"
        with mapping_path.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows(mapping_rows)
        unresolved_path = package_dir / "unresolved-comments.json"
        unresolved_path.write_text(_json(unresolved), encoding="utf-8")
        (package_dir / "response-review-history.json").write_text(_json({
            "scientific_acceptance": False, "automatic_response_approval": False,
            "comments": response_history,
        }), encoding="utf-8")
        artifacts = [
            item
            for item in generated_version["artifacts"]
            if item["artifact_kind"] in {"PDF", "TEX", "BIB", "MARKDOWN", "GENERATED_JSON"}
            and item["document_revision"] == workspace["last_generated_revision"]
        ]
        if not {"PDF", "TEX", "BIB", "MARKDOWN"}.issubset({item["artifact_kind"] for item in artifacts}):
            raise RSTV4Error("The fixed generated version is missing required submission artifacts")
        for artifact in artifacts:
            raw = self._verified_artifact_bytes(artifact)
            clean_name = {
                "PDF": "clean-manuscript.pdf", "TEX": "clean-manuscript.tex",
                "BIB": "references.bib", "MARKDOWN": "clean-manuscript.md",
                "GENERATED_JSON": "generated.json",
            }[artifact["artifact_kind"]]
            (package_dir / clean_name).write_bytes(raw)
        for block in generated_version["blocks"]:
            for asset in block.get("assets", []):
                for representation in ["source", "render"]:
                    key = "source_file" if representation == "source" else "render_file"
                    if key in asset:
                        raw, _, _ = self.generated_material_content(
                            job["manuscript_id"], generated_version["generated_version_id"],
                            asset["asset_id"], representation
                        )
                        (package_dir / asset[key]).write_bytes(raw)
        current_markdown = package_dir / "clean-manuscript.md"
        current_text = current_markdown.read_text(encoding="utf-8")
        if current_text != generated_version["content"]:
            raise RSTV4Error("The clean manuscript does not match its fixed generated version")
        diff_lines = list(difflib.unified_diff(
            baseline.splitlines(),
            current_text.splitlines(),
            fromfile="submitted-baseline.md",
            tofile="revised-manuscript.md",
            lineterm="",
        ))
        (package_dir / "marked-comparison.md").write_text(
            "# Marked manuscript comparison\n\n```diff\n"
            + "\n".join(diff_lines)
            + "\n```\n",
            encoding="utf-8",
        )
        baseline_version = (self.get_generated_version(job["manuscript_id"], review_round["baseline_manuscript_version_id"])
                            if review_round and review_round["baseline_manuscript_version_id"] else None)
        original_document = (self.get_document_revision(job["manuscript_id"], baseline_version["input_revision"])["document"]
                             if baseline_version and baseline_version["input_revision"] is not None else None)
        location_summary = write_mapping(
            package_dir / "change-location-map.csv",
            comments, target=generated_version, baseline=baseline_version,
            target_document=self.get_document_revision(job["manuscript_id"], generated_version["input_revision"])["document"],
            baseline_document=original_document, pdf=package_dir / "clean-manuscript.pdf",
        )
        (package_dir / "author-structure.json").write_text(_json({
            "manuscript_id": job["manuscript_id"],
            "input_revision": generated_version["input_revision"],
            "document": self.get_document_revision(job["manuscript_id"], generated_version["input_revision"])["document"],
        }), encoding="utf-8")
        generated_job_ids = {item["job_id"] for item in generated_version["artifacts"]}
        model_calls = [call for call in self.list_model_calls(job["manuscript_id"])
                       if call["job_id"] in generated_job_ids]
        (package_dir / "generation-model-calls.json").write_text(_json({
            "generated_version_id": generated_version["generated_version_id"],
            "calls": model_calls,
            "provenance_status": "RECORDED_ACTUAL_CALLS" if model_calls else "LEGACY_CALLS_UNKNOWN",
        }), encoding="utf-8")
        manifest = {
            "manuscript_id": job["manuscript_id"],
            "generated_version_id": workspace["active_generated_version_id"],
            "review_round_id": review_round["review_round_id"] if review_round else None,
            "baseline_manuscript_version_id": review_round["baseline_manuscript_version_id"] if review_round else None,
            "baseline_content_sha256": content_digest(baseline) if review_round else None,
            "location_summary": location_summary,
            "input_revision": generated_version["input_revision"],
            "files": sorted({path.name for path in package_dir.iterdir() if path.is_file()} | {"manifest.json"}),
            "file_integrity": file_inventory(package_dir),
            "unresolved_comment_ids": unresolved,
        }
        (package_dir / "manifest.json").write_text(_json(manifest), encoding="utf-8")
        package_id = new_id("submission_package")
        with transaction(self.db_path) as conn:
            conn.execute("BEGIN IMMEDIATE")
            current_revision = conn.execute(
                "SELECT revision FROM studio_documents WHERE manuscript_id=?",
                (job["manuscript_id"],),
            ).fetchone()["revision"]
            if current_revision != job["input_revision"]:
                raise RevisionConflictError("Submission package input changed during packaging")
            if self._package_response_snapshot(conn, job["manuscript_id"], review_round) != response_snapshot:
                raise RevisionConflictError("Replies, associations or author reviews changed during packaging")
            conn.execute(
                "INSERT INTO studio_submission_packages VALUES(?,?,?,?,?,?)",
                (
                    package_id,
                    job["manuscript_id"],
                    review_round["review_round_id"] if review_round else None,
                    job["job_id"],
                    _json(manifest),
                    _now(),
                ),
            )
        return {"package_id": package_id, "directory": str(package_dir), "manifest": manifest}

    def _package_response_snapshot(self, conn: Any, manuscript_id: str, review_round: Any) -> dict[str, Any]:
        if review_round is None:
            return {}
        return {row["comment_id"]: {"state_sha256": response_checks.digest(response_checks.state(conn, manuscript_id, row["comment_id"])),
                                    "status": row["status"],
                                    "content_sha256": response_checks.content_digest(row, response_checks.response_revision(conn, row["comment_id"])),
                                    "verification": response_checks.status(conn, manuscript_id, row)}
                for row in conn.execute("SELECT * FROM review_comments WHERE review_round_id=? AND archived_at IS NULL ORDER BY comment_id",
                                        (review_round["review_round_id"],))}

    def list_submission_packages(self, manuscript_id: str) -> list[dict[str, Any]]:
        return list_packages(self.db_path, self.storage_root, manuscript_id)

    def submission_package_content(self, manuscript_id: str, package_id: str,
                                   filename: str | None = None) -> tuple[bytes, str, str]:
        return package_content(self.db_path, self.storage_root, manuscript_id, package_id, filename)

    def _compile_artifacts(self, job: dict[str, Any], generated: dict[str, Any]) -> list[dict[str, Any]]:
        artifact_dir = self.storage_root / job["manuscript_id"] / "generated" / job["job_id"]
        artifact_dir.mkdir(parents=True, exist_ok=True)
        if generated.get("material_snapshot_schema") == 1:
            verify_material_files(artifact_dir, [
                a for s in generated["sections"] for p in s["paragraphs"] for a in p.get("assets", [])
            ])
        generated["citation_style"] = "numeric_first_use"
        if "references" not in generated:
            selected = [
                evidence_id for section in generated["sections"]
                for paragraph in section["paragraphs"] for evidence_id in paragraph["evidence_ids"]
            ]
            generated["references"] = self._reference_snapshot(generated, self._evidence_by_ids(selected))
        markdown = self._generated_markdown(generated)
        markdown_path = artifact_dir / "manuscript.md"
        markdown_path.write_text(markdown, encoding="utf-8")
        json_path = artifact_dir / "generated.json"
        json_path.write_text(_json(generated), encoding="utf-8")
        evidence = generated["references"]
        bib_path = artifact_dir / "references.bib"
        bib_path.write_text(self._bibtex(evidence), encoding="utf-8")
        tex_path = artifact_dir / "manuscript.tex"
        tex_path.write_text(self._latex(generated, evidence), encoding="utf-8")
        xelatex = shutil.which("xelatex")
        bibtex = shutil.which("bibtex")
        if not xelatex or not bibtex:
            raise RSTV4Error("TeX Live xelatex and bibtex are required")
        env = {**os.environ, "openin_any": "p", "openout_any": "p"}
        latex_command = [
            xelatex,
            "-no-shell-escape",
            "-interaction=nonstopmode",
            "-halt-on-error",
            "manuscript.tex",
        ]
        commands = [latex_command]
        if evidence:
            commands.extend([[bibtex, "manuscript"], latex_command, latex_command])
        else:
            commands.append(latex_command)
        long_windows_input = os.name == "nt" and any(
            len(str(path)) > 240 for path in artifact_dir.iterdir() if path.is_file()
        )
        compile_context = (tempfile.TemporaryDirectory(prefix="rstv4-tex-")
                           if long_windows_input else nullcontext(str(artifact_dir)))
        with compile_context as compile_path:
            working = Path(compile_path)
            if long_windows_input:
                for path in artifact_dir.iterdir():
                    if path.is_file():
                        shutil.copyfile(path, working / path.name)
            for command in commands:
                completed = subprocess.run(
                    command,
                    cwd=working,
                    env=env,
                    capture_output=True,
                    text=True,
                    timeout=90,
                    check=False,
                )
                if completed.returncode != 0:
                    if long_windows_input and (working / "manuscript.log").is_file():
                        shutil.copyfile(working / "manuscript.log", artifact_dir / "manuscript.log")
                    tail = (completed.stdout + "\n" + completed.stderr)[-4000:]
                    raise RSTV4Error(f"LaTeX compilation failed: {tail}")
            if long_windows_input:
                for name in ["manuscript.pdf", "manuscript.log", "manuscript.aux",
                             "manuscript.out", "manuscript.bbl", "manuscript.blg"]:
                    if (working / name).is_file():
                        shutil.copyfile(working / name, artifact_dir / name)
        paths = {
            "MARKDOWN": markdown_path,
            "GENERATED_JSON": json_path,
            "TEX": tex_path,
            "BIB": bib_path,
            "PDF": artifact_dir / "manuscript.pdf",
        }
        saved = []
        with transaction(self.db_path) as conn:
            for kind, path in paths.items():
                content = path.read_bytes()
                artifact_id = new_id("artifact")
                conn.execute(
                    "INSERT INTO studio_artifacts VALUES(?,?,?,?,?,?,?,?,?)",
                    (
                        artifact_id,
                        job["job_id"],
                        job["manuscript_id"],
                        job["input_revision"],
                        kind,
                        str(path),
                        hashlib.sha256(content).hexdigest(),
                        len(content),
                        _now(),
                    ),
                )
                saved.append({"artifact_id": artifact_id, "kind": kind, "path": str(path)})
        return saved

    @staticmethod
    def _generated_markdown(generated: dict[str, Any]) -> str:
        return generated_markdown(generated)

    @staticmethod
    def _reference_snapshot(
        generated: dict[str, Any], evidence: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        by_id = {item["evidence_id"]: item for item in evidence}
        ids = citation_ids(StudioService._generated_markdown(generated))
        if any(item not in by_id for item in ids):
            raise RSTV4Error("Generated text contains an unavailable or unselected citation")
        if any(not re.fullmatch(r"[A-Za-z0-9_.:-]+", item) for item in ids):
            raise RSTV4Error("A citation has an unsafe bibliography key")
        return [
            {**by_id[item], "citation_index": index, "source_snapshot_available": True,
             "reference_snapshot_source": "generation_input"}
            for index, item in enumerate(ids, 1)
        ]

    @staticmethod
    def _latex(generated: dict[str, Any], evidence: list[dict[str, Any]]) -> str:
        known = {item["evidence_id"] for item in evidence}

        def cited_text(raw: str) -> str:
            parts = re.split(r"(\[@[^\]]+\])", raw)
            output = []
            for part in parts:
                match = re.fullmatch(r"\[@([^\]]+)\]", part)
                if match:
                    key = match[1]
                    if key not in known or not re.fullmatch(r"[A-Za-z0-9_.:-]+", key):
                        raise RSTV4Error("Citation does not match the frozen bibliography")
                    output.append(rf"\citep{{{key}}}")
                else:
                    output.append(_escape_latex(part))
            return "".join(output)

        parts = [
            r"\documentclass[11pt]{article}",
            r"\usepackage[margin=1in]{geometry}",
            r"\usepackage{fontspec}",
            r"\usepackage{hyperref}",
            r"\usepackage[numbers,square]{natbib}",
            r"\usepackage{amsmath}",
            *([r"\usepackage{graphicx}", r"\usepackage{longtable}"]
              if generated.get("material_snapshot_schema") == 1 else []),
            r"\setmainfont{TeX Gyre Termes}",
            r"\begin{document}",
            rf"\title{{{cited_text(generated['title'])}}}",
            r"\maketitle",
        ]
        if generated.get("abstract"):
            parts.extend([
                r"\begin{abstract}",
                cited_text(generated["abstract"]),
                r"\end{abstract}",
            ])
        for section in generated["sections"]:
            command = ["section", "subsection", "subsubsection", "paragraph", "subparagraph"][
                section.get("section_depth", 0)
            ]
            parts.append("\\" + command + "{" + cited_text(section["title"]) + "}")
            for paragraph in section["paragraphs"]:
                text = cited_text(paragraph["text"])
                parts.extend([text, ""])
                for formula in paragraph.get("formulas", []):
                    StudioService._validate_formula(str(formula.get("latex") or ""))
                    parts.extend([
                        r"\begin{equation}",
                        str(formula["latex"]),
                        r"\end{equation}",
                    ])
                    if formula.get("caption"):
                        parts.extend([cited_text(str(formula["caption"])), ""])
                parts.extend(material_latex(paragraph.get("assets", []), _escape_latex))
        if evidence:
            parts.extend([r"\bibliographystyle{unsrtnat}", r"\bibliography{references}"])
        parts.append(r"\end{document}")
        return "\n".join(parts) + "\n"

    @staticmethod
    def _bibtex(evidence: list[dict[str, Any]]) -> str:
        entries = []
        for item in evidence:
            authors = " and ".join(item.get("authors") or ["Unknown"])
            fields = {
                "title": item.get("title") or item.get("statement") or "Untitled source",
                "author": authors,
                "year": str(item.get("published_year") or "n.d."),
            }
            if item.get("doi"):
                fields["doi"] = item["doi"]
            body = ",\n".join(
                f"  {key} = {{{_escape_latex(str(value))}}}" for key, value in fields.items()
            )
            entries.append(f"@article{{{item['evidence_id']},\n{body}\n}}")
        return "\n\n".join(entries) + ("\n" if entries else "")

    def _evidence_by_ids(self, evidence_ids: list[str]) -> list[dict[str, Any]]:
        unique = list(dict.fromkeys(evidence_ids))
        if not unique:
            return []
        placeholders = ",".join("?" for _ in unique)
        with transaction(self.db_path) as conn:
            rows = conn.execute(
                f"""SELECT e.*,p.title,p.authors_json,p.published_year,p.doi,p.landing_url,p.pdf_url
                    FROM evidence e LEFT JOIN papers p ON p.paper_id=e.paper_id
                    WHERE e.evidence_id IN ({placeholders})""",
                unique,
            ).fetchall()
            items = []
            for row in rows:
                item = dict(row)
                item["authors"] = json.loads(item.pop("authors_json")) if item["authors_json"] else []
                item["location"] = json.loads(item.pop("location_json"))
                item["source_context"] = freeze_source(conn, item)
                items.append(item)
        return items

    @staticmethod
    def _validate_svg(content: bytes) -> dict[str, Any]:
        if len(content) > 5 * 1024 * 1024:
            raise ValueError("SVG exceeds the 5 MB safe parsing limit")
        try:
            root = ET.fromstring(content)
        except ET.ParseError as exc:
            raise ValueError("Invalid SVG") from exc
        forbidden = {"script", "foreignObject", "iframe", "object", "embed"}
        for element in root.iter():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag in forbidden:
                raise ValueError(f"Unsafe SVG element: {tag}")
            for key, value in element.attrib.items():
                local = key.rsplit("}", 1)[-1].casefold()
                lowered = value.strip().casefold()
                if local.startswith("on") or (
                    local in {"href", "src"}
                    and not lowered.startswith("#")
                ):
                    raise ValueError("SVG contains an unsafe external reference")
        return {"root": root.tag.rsplit("}", 1)[-1], "safe": True}

    @staticmethod
    def _preview_csv(content: bytes) -> dict[str, Any]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV must be UTF-8") from exc
        rows = list(csv.reader(io.StringIO(text)))
        return {"rows": len(rows), "columns": max((len(row) for row in rows), default=0), "preview": rows[:6]}

    @staticmethod
    def _preview_markdown_table(content: bytes) -> dict[str, Any]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Markdown tables must be UTF-8") from exc
        lines = [line.strip() for line in text.splitlines() if "|" in line][:100]
        preview = [
            [cell.strip() for cell in line.strip("|").split("|")][:30]
            for line in lines
            if not re.fullmatch(r"[|:\-\s]+", line)
        ]
        return {
            "rows": len(preview),
            "columns": max((len(row) for row in preview), default=0),
            "preview": preview[:6],
        }

    @staticmethod
    def _preview_xlsx(content: bytes) -> dict[str, Any]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                total = sum(info.file_size for info in archive.infolist())
                if total > 100 * 1024 * 1024:
                    raise ValueError("XLSX expands beyond the 100 MB safety limit")
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid XLSX workbook") from exc
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            preview = []
            for row in sheet.iter_rows(max_row=6, max_col=6):
                values = [cell.value for cell in row]
                while values and values[-1] is None:
                    values.pop()
                preview.append(values)
            return {
                "sheets": workbook.sheetnames,
                "active_sheet": sheet.title,
                "preview": preview,
                "formula_policy": "cached-values-only",
            }
        finally:
            workbook.close()

    @staticmethod
    def _validate_pdf_asset(content: bytes) -> dict[str, Any]:
        if not content.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF signature")
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF assets are not supported")
        return {"pages": len(reader.pages), "safe": True}

    @staticmethod
    def _validate_docx_asset(content: bytes) -> dict[str, Any]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                expanded = sum(item.file_size for item in members)
                if len(members) > 10_000 or expanded > 100 * 1024 * 1024:
                    raise ValueError("DOCX expands beyond the safe import limit")
                if "[Content_Types].xml" not in archive.namelist():
                    raise ValueError("Invalid DOCX package")
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid DOCX package") from exc
        return {"expanded_bytes": expanded, "safe": True}

    @staticmethod
    def _validate_formula(latex: str) -> None:
        if not latex.strip():
            raise ValueError("Formula LaTeX is empty")
        if len(latex) > 12_000:
            raise ValueError("Formula LaTeX exceeds the 12,000 character limit")
        forbidden = re.compile(
            r"\\(?:input|include|openin|openout|write|read|usepackage|documentclass|"
            r"catcode|csname|newcommand|renewcommand|def|loop|repeat|special|immediate|"
            r"shellesc)\b",
            re.IGNORECASE,
        )
        if forbidden.search(latex):
            raise ValueError("Formula contains a forbidden TeX command")

    def _job_cancel_requested(self, job_id: str) -> bool:
        with transaction(self.db_path) as conn:
            row = conn.execute(
                "SELECT status FROM studio_jobs WHERE job_id=?", (job_id,)
            ).fetchone()
        return row is None or row["status"] == "CANCEL_REQUESTED"

    @staticmethod
    def _saved_response(row: Any, operation: dict[str, Any], *, idempotent: bool) -> dict[str, Any]:
        return {
            "document_revision": row["revision"],
            "document": json.loads(row["document_json"]),
            "operation": operation,
            "idempotent_replay": idempotent,
            "saved_at": row["updated_at"],
        }

    @staticmethod
    def _asset(row: Any) -> dict[str, Any]:
        return _decode(row, "metadata_json")

    @staticmethod
    def _artifact(row: Any) -> dict[str, Any]:
        return dict(row)

    @staticmethod
    def _message(row: Any) -> dict[str, Any]:
        result = _decode(row, "quoted_json")
        provenance = result["quoted"].pop("_model_provenance", None)
        if provenance is not None:
            result["model_provenance"] = provenance
        return result

    @staticmethod
    def _comment(row: Any) -> dict[str, Any]:
        result = _decode(row, "import_metadata_json")
        number = result["import_metadata"].get("comment_number")
        result["display_label"] = f"A{number:02d}" if number else result["comment_id"]
        return result

    @staticmethod
    def _finding(row: Any) -> dict[str, Any]:
        return _decode(row, "anchor_json")

    @staticmethod
    def _job(row: Any) -> dict[str, Any]:
        result = dict(row)
        for field in ("request_json", "result_json", "error_json"):
            value = result.pop(field)
            result[field.removesuffix("_json")] = json.loads(value) if value else None
        return result
