from pathlib import Path
from io import BytesIO, StringIO
import csv
import hashlib
import json
import os
import shutil
import sqlite3
import subprocess
import tempfile
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader

r=Path.cwd()
fixture=json.loads((r/'.tmp/scholar-generated-assets-native-v1.json').read_text(encoding='utf-8'))
ms=fixture['manuscript_id']
root=Path(fixture['root'])
out=r/'.tmp/scholar-generated-assets-native-audit-v1'
out.mkdir()
def digest(raw):return hashlib.sha256(raw).hexdigest()
def snapshot():
    with sqlite3.connect(f"file:{(root/'data/catalog.db').as_posix()}?mode=ro",uri=True) as c:
        c.row_factory=sqlite3.Row
        return {x[0]:[dict(row) for row in c.execute('SELECT * FROM "'+x[0]+'" ORDER BY rowid')]
                for x in c.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")}
now=snapshot()
prior=fixture['after_fixture_tables']
assert now.keys()==prior.keys() and len(now)==109
allowed={'manuscripts':{'writing_status','revision','updated_at','workflow_stage'},
         'studio_documents':{'active_generated_version_id','last_generated_revision','updated_at'}}
diffs=[]
for table,rows in prior.items():
    for row in rows:
        if row in now[table]:continue
        assert table in allowed and row.get('manuscript_id')==ms,(table,row)
        current=[v for v in now[table] if v.get('manuscript_id')==ms]
        assert len(current)==1
        changed={key for key in row if row[key]!=current[0][key]}
        assert changed<=allowed[table],(table,changed)
        diffs.append({'table':table,'manuscript_id':ms,'changed_columns':sorted(changed)})
assert now['approval_grants']==prior['approval_grants']
doc=next(x for x in now['studio_documents'] if x['manuscript_id']==ms)
assert json.loads(doc['document_json'])==fixture['initial_document'] and doc['revision']==1
manuscript=next(x for x in now['manuscripts'] if x['manuscript_id']==ms)
assert manuscript['workflow_stage']=='FINALIZATION_SUBMISSION' and manuscript['revision']==3
jobs=[x for x in now['studio_jobs'] if x['manuscript_id']==ms]
assert len(jobs)==2 and {j['job_kind'] for j in jobs}=={'GENERATE','SUBMISSION_PACKAGE'}
assert all(j['status']=='SUCCEEDED' and j['error_json'] is None and j['input_revision']==1 for j in jobs)
gen_job=next(j for j in jobs if j['job_kind']=='GENERATE')
pkg_job=next(j for j in jobs if j['job_kind']=='SUBMISSION_PACKAGE')
gen_result=json.loads(gen_job['result_json'])
pkg_result=json.loads(pkg_job['result_json'])
version_id=gen_result['generated_version_id']
assert doc['active_generated_version_id']==version_id and doc['last_generated_revision']==1
artifacts=[a for a in now['studio_artifacts'] if a['job_id']==gen_job['job_id']]
assert len(artifacts)==5 and {a['artifact_kind'] for a in artifacts}=={'PDF','MARKDOWN','GENERATED_JSON','BIB','TEX'}
for artifact in artifacts:
    raw=Path(artifact['storage_path']).read_bytes()
    assert len(raw)==artifact['size_bytes'] and digest(raw)==artifact['sha256']
generated=json.loads(Path(next(a['storage_path'] for a in artifacts if a['artifact_kind']=='GENERATED_JSON')).read_text(encoding='utf-8'))
paragraph=generated['sections'][0]['paragraphs'][0]
assets=paragraph['assets']
assert generated['material_snapshot_schema']==1
assert paragraph['asset_ids']==fixture['initial_document']['sections'][0]['paragraphs'][0]['asset_ids']
assert [a['asset_id'] for a in assets]==paragraph['asset_ids'] and len(assets)==6
assert paragraph['text']==fixture['initial_document']['sections'][0]['paragraphs'][0]['intents'][0]['thoughts'][0]['text']
calls=[c for c in now['studio_model_calls'] if c['manuscript_id']==ms]
assert len(calls)==1 and calls[0]['status']=='DONE'
assert json.loads(calls[0]['contract_json'])['actual_model']=='TEST_FIXTURE_MATERIALS_NATIVE'
assert json.loads(calls[0]['input_json'])['assets']==assets
visible=json.loads((r/'.tmp/scholar-generated-assets-native-proof-v1/generated-visible-cells-v1.json').read_text(encoding='utf-8'))
reloaded=json.loads((r/'.tmp/scholar-generated-assets-native-proof-v1/generated-after-refresh-v1.json').read_text(encoding='utf-8'))
assert visible['tables']==reloaded['tables'] and visible['sources']==reloaded['sources']
assert len(visible['images'])==3 and all(v['complete'] and v['width']>0 and v['height']>0 for v in visible['images'])
assert all(version_id in v['src'] for v in visible['images'])
assert all(version_id in v['url'] for v in visible['sources'])
full_tables=[{'title':t['title'],'rows':t['rows']} for a in assets for t in a.get('tables',[])]
assert full_tables==visible['tables'] and len(full_tables)==4
material_receipts=[]
directory=Path(next(a['storage_path'] for a in artifacts if a['artifact_kind']=='PDF')).parent
for asset,original in zip(assets,fixture['assets'],strict=True):
    raw=Path(original['storage_path']).read_bytes()
    assert asset['sha256']==digest(raw)==original['sha256'] and len(raw)==asset['size_bytes']
    assert (directory/asset['source_file']).read_bytes()==raw
    if asset['filename'].endswith('.csv'):
        assert list(csv.reader(StringIO(raw.decode('utf-8'))))==asset['tables'][0]['rows']
    if asset['filename'].endswith('.xlsx'):
        book=load_workbook(BytesIO(raw),data_only=True)
        assert [s.title for s in book]==[t['title'] for t in asset['tables']]
        assert [[[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(values_only=True)] for sheet in book]==[t['rows'] for t in asset['tables']]
        book.close()
    if asset['filename'].endswith('.md'):
        original_rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in raw.decode('utf-8').splitlines()]
        assert original_rows[:1]+original_rows[2:]==asset['tables'][0]['rows']
    if asset.get('render_file'):
        rendered=(directory/asset['render_file']).read_bytes()
        assert digest(rendered)==asset['render_sha256'] and len(rendered)==asset['render_size_bytes']
    material_receipts.append({k:asset[k] for k in ['asset_id','filename','media_type','sha256','size_bytes','source_file']})
before_downloads=json.loads((r/'.tmp/scholar-generated-assets-downloads-before-v1.json').read_text(encoding='utf-8'))
downloads=Path.home()/'Downloads'
for name,data in before_downloads.items():
    assert (downloads/name).stat().st_size==data['bytes'] and digest((downloads/name).read_bytes())==data['sha256']
pdf_artifact=next(a for a in artifacts if a['artifact_kind']=='PDF')
pdfs=[p for p in downloads.glob('manuscript*.pdf') if p.name not in before_downloads and digest(p.read_bytes())==pdf_artifact['sha256']]
assert len(pdfs)==1,'Native downloaded PDF missing or ambiguous'
pdf=pdfs[0]
zip_path=downloads/(pkg_result['package_id']+'.zip')
assert zip_path.exists() and zip_path.stat().st_size>0
pkg_manifest=pkg_result['manifest']
assert pkg_manifest['manuscript_id']==ms and pkg_manifest['generated_version_id']==version_id and pkg_manifest['input_revision']==1
assert pkg_manifest['review_round_id'] is None and pkg_manifest['baseline_manuscript_version_id'] is None
independent=Path(tempfile.mkdtemp(prefix='sc-native-material-package-')).resolve()
assert independent.parent==Path(tempfile.gettempdir()).resolve()
with zipfile.ZipFile(zip_path) as archive:
    assert archive.testzip() is None
    assert len(archive.namelist())==len(set(archive.namelist()))==25
    assert sorted(archive.namelist())==sorted(pkg_manifest['files'])
    assert json.loads(archive.read('manifest.json'))==pkg_manifest
    for name,record in pkg_manifest['file_integrity'].items():
        raw=archive.read(name)
        assert len(raw)==record['size_bytes'] and digest(raw)==record['sha256']
    for asset,original in zip(assets,fixture['assets'],strict=True):
        assert archive.read(asset['source_file'])==Path(original['storage_path']).read_bytes()
        if asset.get('render_file'):assert digest(archive.read(asset['render_file']))==asset['render_sha256']
    assert archive.read('clean-manuscript.pdf')==pdf.read_bytes()
    for name in archive.namelist():
        assert Path(name).name==name and (independent/name).resolve().is_relative_to(independent)
    archive.extractall(independent)
result=subprocess.run([shutil.which('xelatex'),'-no-shell-escape','-interaction=nonstopmode','-halt-on-error','clean-manuscript.tex'],cwd=independent,capture_output=True,text=True,encoding='utf-8',errors='replace')
(out/'independent-recompile.log').write_text(result.stdout+result.stderr,encoding='utf-8')
assert result.returncode==0 and (independent/'clean-manuscript.pdf').stat().st_size>0
pages=[]
reader=PdfReader(pdf)
for i,page in enumerate(reader.pages,1):
    text=page.extract_text() or ''
    (out/f'page-{i}-text.txt').write_text(text,encoding='utf-8')
    pages.append({'page':i,'material_labels':[f'{kind}_ROW_{n}' for kind in ['CSV','XLSX','MD'] for n in range(1,11) if f'{kind}_ROW_{n}' in text],
                  'second_sheet':'SECOND_SHEET_ORIGINAL' in text,'images':[{'name':image.name,'size':image.image.size} for image in page.images]})
text='\n'.join(p.extract_text() or '' for p in reader.pages)
assert all(f'{kind}_ROW_{n}' in text for kind in ['CSV','XLSX','MD'] for n in range(1,11))
assert 'SECOND_SHEET_ORIGINAL' in text and paragraph['text'] in text
assert sum(len(p['images']) for p in pages)==3
assert any(all(f'CSV_ROW_{n}' in p['material_labels'] for n in range(1,11)) for p in pages),'Compact CSV should remain together'
render_dir=out/'render'
render_dir.mkdir()
renderer=shutil.which('pdftoppm')
assert renderer
render_result=subprocess.run([renderer,'-png','-r','110',str(pdf),str(render_dir/'page')],capture_output=True,text=True)
assert render_result.returncode==0,render_result.stderr
pngs=sorted(render_dir.glob('page-*.png'))
assert len(pngs)==len(reader.pages) and digest(pdf.read_bytes())==pdf_artifact['sha256']
assert snapshot()==now
(out/'all109tables-after-native.json').write_text(json.dumps(now,ensure_ascii=False,indent=2),encoding='utf-8')
receipt={'status':'PASS_NATIVE_MATERIAL_DELIVERY_AND_INDEPENDENT_AUDIT','manuscript_id':ms,'generated_version_id':version_id,
         'author_revision':1,'stage':'FINALIZATION_SUBMISSION','job_ids':[j['job_id'] for j in jobs],
         'declared_fixture_model_calls':1,'actual_remote_model_calls':0,'artifacts':artifacts,
         'materials':material_receipts,'full_tables':full_tables,'source_and_browser_rows_exact':True,
         'all109tables_preserve_foreign_records':True,'owned_row_diffs':diffs,'downloaded_pdf':str(pdf),
         'pdf_sha256':digest(pdf.read_bytes()),'downloaded_zip':str(zip_path),'zip_sha256':digest(zip_path.read_bytes()),
         'zip_bytes':zip_path.stat().st_size,'zip_files':25,'integrity_inventory_files':len(pkg_manifest['file_integrity']),
         'package_id':pkg_result['package_id'],'package_directory':pkg_result['directory'],
         'independent_recompile_directory':str(independent),'independent_recompile_exit':result.returncode,
         'pdf_pages':pages,'rendered_pages':[str(p) for p in pngs],'renderer':renderer,
         'pdf_visual_acceptance':'PENDING_ACTUAL_IMAGE_INSPECTION','scientific_acceptance':False,
         'original118_complete':False,'all11_complete':False,'download_testing_resumed':True}
(out/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:receipt[k] for k in ['status','downloaded_pdf','pdf_sha256','downloaded_zip','zip_bytes','zip_files','independent_recompile_exit','rendered_pages','pdf_visual_acceptance','all109tables_preserve_foreign_records']}))
