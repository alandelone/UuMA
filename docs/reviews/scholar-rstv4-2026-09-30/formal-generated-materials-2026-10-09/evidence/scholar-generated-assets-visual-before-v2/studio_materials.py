"""Freeze original paragraph materials and render only their verified, complete content."""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import sqlite3
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from .service import RSTV4Error

MAX_CELLS = 20000
MAX_PIXELS = 40000000
SNAPSHOT_SCHEMA = 1


def digest(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def checked_svg(raw: bytes) -> str:
    text = raw.decode("utf-8-sig")
    if len(raw) > 5 * 1024 * 1024 or re.search(
        r"<!DOCTYPE|<!ENTITY|@import", text, re.IGNORECASE
    ):
        raise ValueError("SVG contains unsupported external declarations")
    root = ET.fromstring(text)
    if root.tag.rsplit("}", 1)[-1] != "svg":
        raise ValueError("SVG must have an svg root")

    def extent(value):
        match = re.fullmatch(r"\s*([0-9.+eE-]+)\s*(px|pt|pc|mm|cm|in)?\s*", value or "")
        if not match:
            return None
        number = float(match[1])
        factor = {
            None: 1,
            "px": 1,
            "pt": 96 / 72,
            "pc": 16,
            "mm": 96 / 25.4,
            "cm": 96 / 2.54,
            "in": 96,
        }[match[2]]
        return number * factor if math.isfinite(number) and number > 0 else None

    width, height = extent(root.get("width")), extent(root.get("height"))
    if width is None or height is None:
        view = [
            float(v) for v in re.split(r"[\s,]+", root.get("viewBox", "").strip()) if v
        ]
        if (
            len(view) != 4
            or not all(math.isfinite(v) for v in view)
            or view[2] <= 0
            or view[3] <= 0
        ):
            raise ValueError("SVG needs finite original dimensions or a viewBox")
        width, height = view[2:]
    if 1800 * math.ceil(1800 * height / width) > MAX_PIXELS:
        raise ValueError("SVG exceeds the complete render pixel limit")
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] in {
            "script",
            "foreignObject",
            "iframe",
            "object",
            "embed",
        }:
            raise ValueError("SVG contains executable content")
        for key, value in element.attrib.items():
            local = key.rsplit("}", 1)[-1].casefold()
            if local.startswith("on") or (
                local in {"href", "src"} and not value.strip().startswith("#")
            ):
                raise ValueError("SVG contains an external reference")
        for value in [*element.attrib.values(), element.text or ""]:
            for target in re.findall(r"url\s*\((.*?)\)", value, re.IGNORECASE):
                if not target.strip().strip("'\"").startswith("#"):
                    raise ValueError("SVG contains an external resource")
    return text


def _table(rows: list[list[Any]], title: str) -> dict[str, Any]:
    if not rows or not any(
        any(value not in (None, "") for value in row) for row in rows
    ):
        raise ValueError("Table has no original cell data")
    columns = max(map(len, rows))
    if len(rows) * columns > MAX_CELLS or columns > 40:
        raise ValueError(
            "Table exceeds the complete-render limit (20000 cells / 40 columns)"
        )
    return {
        "title": title,
        "rows": [[_cell(value) for value in row] for row in rows],
        "row_count": len(rows),
        "column_count": columns,
    }


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _markdown_rows(text: str) -> list[list[str]]:
    rows = []
    for line in text.splitlines():
        if not line.strip():
            continue
        if "|" not in line:
            raise ValueError(
                "Markdown material must be a complete table; extra prose is not discarded"
            )
        cells = re.split(r"(?<!\\)\|", line.strip())
        if not cells[0]:
            cells.pop(0)
        if cells and not cells[-1]:
            cells.pop()
        values = [cell.strip().replace(r"\|", "|") for cell in cells]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in values):
            continue
        rows.append(values)
    return rows


def representation(
    row: dict[str, Any], raw: bytes
) -> tuple[dict[str, Any], bytes | None]:
    mime = row["media_type"]
    if mime in {"image/png", "image/jpeg"}:
        from PIL import Image, ImageOps, __version__

        with Image.open(io.BytesIO(raw)) as image:
            expected = "PNG" if mime == "image/png" else "JPEG"
            if image.format != expected or getattr(image, "n_frames", 1) != 1:
                raise ValueError(
                    "Image type or frame count does not match its original material"
                )
            if image.width * image.height > MAX_PIXELS:
                raise ValueError("Image exceeds the 40 million pixel render limit")
            image.load()
            if mime == "image/png" and image.getexif().get(274, 1) == 1:
                return {
                    "display_kind": "FIGURE",
                    "width": image.width,
                    "height": image.height,
                    "renderer": f"Original PNG verified with Pillow/{__version__}",
                }, raw
            normalized = ImageOps.exif_transpose(image)
            if normalized.mode not in {"RGB", "RGBA", "L", "LA"}:
                normalized = normalized.convert(
                    "RGBA" if "transparency" in normalized.info else "RGB"
                )
            buf = io.BytesIO()
            normalized.save(buf, format="PNG")
            return {
                "display_kind": "FIGURE",
                "width": normalized.width,
                "height": normalized.height,
                "renderer": f"Pillow/{__version__}",
            }, buf.getvalue()
    if mime == "image/svg+xml":
        import resvg_py
        from PIL import Image

        text = checked_svg(raw)
        rendered = resvg_py.svg_to_bytes(svg_string=text, width=1800)
        with Image.open(io.BytesIO(rendered)) as image:
            if image.width * image.height > MAX_PIXELS:
                raise ValueError("SVG exceeds the complete render pixel limit")
            return {
                "display_kind": "FIGURE",
                "width": image.width,
                "height": image.height,
                "renderer": f"resvg_py/{resvg_py.__version__}",
            }, rendered
    if mime == "text/csv":
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        return {
            "display_kind": "TABLE",
            "tables": [_table(rows, row["filename"])],
        }, None
    if mime == "text/markdown":
        rows = _markdown_rows(raw.decode("utf-8-sig"))
        return {
            "display_kind": "TABLE",
            "tables": [_table(rows, row["filename"])],
        }, None
    if mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        from openpyxl import load_workbook

        values = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        tables = []
        try:
            for sheet in values:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                    raise ValueError(
                        f"Workbook sheet {sheet.title} exceeds the complete render limit"
                    )
                source = formulas[sheet.title]
                rows = []
                for value_row, formula_row in zip(
                    sheet.iter_rows(), source.iter_rows(), strict=True
                ):
                    row_values = []
                    for cell, original in zip(value_row, formula_row, strict=True):
                        if original.data_type == "f" and cell.value is None:
                            raise ValueError(
                                f"Workbook formula {sheet.title}!{original.coordinate} has no cached original value"
                            )
                        row_values.append(cell.value)
                    rows.append(row_values)
                if any(any(value not in (None, "") for value in row) for row in rows):
                    tables.append(_table(rows, sheet.title))
            if not tables:
                raise ValueError("Workbook has no original cell data")
            if sum(t["row_count"] * t["column_count"] for t in tables) > MAX_CELLS:
                raise ValueError("Workbook exceeds the complete render cell limit")
            return {
                "display_kind": "TABLE",
                "tables": tables,
                "formula_policy": "cached-values-only",
            }, None
        finally:
            values.close()
            formulas.close()
    if row["asset_kind"] == "DOCUMENT":
        return {"display_kind": "ATTACHMENT"}, None
    raise ValueError(f"Unsupported original material: {mime}")


def freeze(
    conn: sqlite3.Connection,
    storage_root: Path,
    manuscript_id: str,
    document: dict[str, Any],
) -> tuple[dict[str, dict], dict[str, bytes]]:
    ids = list(
        dict.fromkeys(
            a
            for s in document["sections"]
            for p in s["paragraphs"]
            for a in p.get("asset_ids", [])
        )
    )
    snapshots, files = {}, {}
    for asset_id in ids:
        if not isinstance(asset_id, str) or not re.fullmatch(
            r"[A-Za-z0-9_-]{1,128}", asset_id
        ):
            raise ValueError("Original material has an invalid identity")
        found = conn.execute(
            "SELECT * FROM studio_assets WHERE manuscript_id=? AND asset_id=? AND status='READY'",
            (manuscript_id, asset_id),
        ).fetchone()
        if found is None:
            raise ValueError(f"Original material is unavailable: {asset_id}")
        row = dict(found)
        path = Path(row["storage_path"]).resolve()
        owner = (storage_root / manuscript_id / "assets").resolve()
        if owner not in path.parents or not path.is_file():
            raise ValueError(f"Original material file is unavailable: {asset_id}")
        raw = path.read_bytes()
        if len(raw) != row["size_bytes"] or digest(raw) != row["sha256"]:
            raise ValueError(
                f"Original material failed its recorded byte check: {asset_id}"
            )
        ext = {
            "image/png": "png",
            "image/jpeg": "jpg",
            "image/svg+xml": "svg",
            "text/csv": "csv",
            "text/markdown": "md",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
            "application/pdf": "pdf",
            "text/plain": "txt",
            "application/x-tex": "tex",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        }[row["media_type"]]
        source_name = f"material-{asset_id}-source.{ext}"
        shown, render = representation(row, raw)
        snapshot = {
            key: row[key]
            for key in [
                "asset_id",
                "filename",
                "media_type",
                "asset_kind",
                "sha256",
                "size_bytes",
            ]
        }
        snapshot.update(shown, source_file=source_name, snapshot_schema=SNAPSHOT_SCHEMA)
        files[source_name] = raw
        if render is not None:
            render_name = f"material-{asset_id}-render.png"
            files[render_name] = render
            snapshot.update(
                render_file=render_name,
                render_sha256=digest(render),
                render_size_bytes=len(render),
            )
        snapshots[asset_id] = snapshot
    return snapshots, files


def verify_files(directory: Path, assets: list[dict]) -> None:
    for asset in assets:
        for key, sha_key, size_key in [
            ("source_file", "sha256", "size_bytes"),
            ("render_file", "render_sha256", "render_size_bytes"),
        ]:
            name = asset.get(key)
            if not name:
                continue
            if not isinstance(name, str) or not re.fullmatch(
                r"material-[A-Za-z0-9_-]+-(source|render)\.[a-z0-9]+", name
            ):
                raise RSTV4Error(
                    "Frozen original material has an invalid file identity"
                )
            path = (directory / name).resolve()
            if path.parent != directory.resolve() or not path.is_file():
                raise RSTV4Error("Frozen original material is unavailable")
            raw = path.read_bytes()
            if len(raw) != asset[size_key] or digest(raw) != asset[sha_key]:
                raise RSTV4Error("Frozen original material failed its integrity check")
        if asset["display_kind"] == "TABLE":
            original = (directory / asset["source_file"]).read_bytes()
            parsed, _ = representation(asset, original)
            if asset["tables"] != parsed["tables"]:
                raise RSTV4Error(
                    "Frozen table differs from its complete original cells"
                )


def markdown(assets: list[dict]) -> list[str]:
    parts = []
    for asset in assets:
        label = asset["filename"].replace("[", r"\[").replace("]", r"\]")
        if asset["display_kind"] == "FIGURE":
            parts.extend(["", f"![{label}]({asset['render_file']})"])
        elif asset["display_kind"] == "TABLE":
            for table in asset["tables"]:
                parts.extend(["", f"**{table['title']}**"])
                rows = table["rows"]
                columns = table["column_count"]

                def line(row, column_count=columns):
                    return (
                        "| "
                        + " | ".join(
                            (row[i] if i < len(row) else "")
                            .replace("&", "&amp;")
                            .replace("<", "&lt;")
                            .replace(">", "&gt;")
                            .replace("[", "&#91;")
                            .replace("]", "&#93;")
                            .replace("|", r"\|")
                            .replace("\n", "<br>")
                            for i in range(column_count)
                        )
                        + " |"
                    )

                parts.append(line(rows[0]))
                parts.append("| " + " | ".join("---" for _ in range(columns)) + " |")
                parts.extend(line(row) for row in rows[1:])
        parts.extend(["", f"[Original material: {label}]({asset['source_file']})"])
    return parts


def latex(assets: list[dict], escape) -> list[str]:
    parts = []
    for asset in assets:
        label = escape(asset["filename"])
        if asset["display_kind"] == "FIGURE":
            parts.extend(
                [
                    r"\par\noindent\begin{minipage}{\linewidth}\centering",
                    rf"\includegraphics[width=\linewidth,height=.65\textheight,keepaspectratio]{{{asset['render_file']}}}",
                    rf"\par\small {label}",
                    r"\end{minipage}\par",
                ]
            )
        elif asset["display_kind"] == "TABLE":
            for table in asset["tables"]:
                columns = table["column_count"]
                compact = table["row_count"] <= 20 and columns <= 8
                environment = "tabular" if compact else "longtable"
                if compact:
                    parts.append(r"\par\noindent\begin{minipage}{\linewidth}")
                parts.extend(
                    [
                        rf"\par\noindent\textbf{{{escape(table['title'])}}}",
                        r"\begingroup\small",
                        r"\setlength{\tabcolsep}{3pt}",
                        "\\begin{"
                        + environment
                        + "}{"
                        + "|"
                        + "|".join([f"p{{{0.9 / columns:.5f}\\linewidth}}"] * columns)
                        + "|}",
                        r"\hline",
                    ]
                )
                if not compact:
                    parts.extend(
                        [
                            r"\endfirsthead",
                            rf"\multicolumn{{{columns}}}{{l}}{{{escape(table['title'])} (continued)}} \\ \hline",
                            r"\endhead",
                        ]
                    )
                for row in table["rows"]:
                    cells = [
                        escape(row[i] if i < len(row) else "").replace(
                            "\n", r"\newline "
                        )
                        for i in range(columns)
                    ]
                    parts.append(" & ".join(cells) + r" \\ \hline")
                parts.extend(["\\end{" + environment + "}", r"\endgroup"])
                if compact:
                    parts.append(r"\end{minipage}\par")
        else:
            parts.append(
                rf"\par\noindent Attachment: {label} (original file in the submission package)."
            )
    return parts
