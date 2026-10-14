"""Full material-output contracts; all data and the model are declared software fixtures."""

import hashlib
import json
import shutil
import sqlite3
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader
from test_writing_studio import complete_document, create_studio, wait_for_job

from rstv4.service import RSTV4Error
from rstv4.studio import DocumentValidationError
from rstv4.studio_materials import checked_svg, representation


def six_materials(service, ms):
    files = []
    for kind, mime in [("PNG", "image/png"), ("JPEG", "image/jpeg")]:
        buf = BytesIO()
        image = Image.new("RGB", (320, 160), (40, 120, 180))
        image.save(buf, format=kind)
        files.append((f"original-{kind.lower()}.{kind.lower()}", mime, buf.getvalue()))
    files.append(
        (
            "original-svg.svg",
            "image/svg+xml",
            b'<svg xmlns="http://www.w3.org/2000/svg" width="320" height="160"><rect width="320" height="160" fill="#d2873b"/><circle cx="160" cy="80" r="50" fill="#467344"/></svg>',
        )
    )
    files.append(
        (
            "original-csv.csv",
            "text/csv",
            (
                "Item,Value,Condition\n"
                + "".join(f"CSV_ROW_{n},{n},scope_{n}\n" for n in range(1, 11))
            ).encode("utf-8"),
        )
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Original first sheet"
    for n in range(1, 11):
        sheet.append([f"XLSX_ROW_{n}", n, f"scope_{n}"])
    second = workbook.create_sheet("Original second sheet")
    second.append(["SECOND_SHEET_ORIGINAL", 7, "Explicit condition"])
    buf = BytesIO()
    workbook.save(buf)
    files.append(
        (
            "original-xlsx.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            buf.getvalue(),
        )
    )
    files.append(
        (
            "original-markdown.md",
            "text/markdown",
            (
                "| Item | Value | Condition |\n|---|---|---|\n"
                + "".join(f"| MD_ROW_{n} | {n} | scope_{n} |\n" for n in range(1, 11))
            ).encode("utf-8"),
        )
    )
    assets = [service.store_asset(ms, name, mime, raw) for name, mime, raw in files]
    return assets, files


def test_all_six_original_materials_appear_in_frozen_view_and_real_pdf(tmp_path):
    service, _, ms = create_studio(tmp_path)
    doc = complete_document(service.get_workspace(ms))
    assets, files = six_materials(service, ms)
    paragraph = doc["sections"][0]["paragraphs"][0]
    paragraph["asset_ids"] = [a["asset_id"] for a in assets]
    service.save_document(
        ms, base_revision=0, operation_id="six-materials", document=doc
    )
    captured = []

    def realize(task, payload):
        captured.append(deepcopy(payload))
        thought = payload["intents"][0]["thoughts"][0]
        return {
            "paragraph": thought["text"],
            "sentence_mappings": [{"sentence": 0, "thought_ids": [thought["id"]]}],
            "_contract": {"actual_model": "TEST_FIXTURE_MATERIALS", "task": task},
        }

    service.gateway.analyze = realize
    try:
        job = wait_for_job(
            service, service.create_job(ms, "GENERATE", input_revision=1)
        )
        assert job["status"] == "SUCCEEDED", job
        version = service.get_generated_version(
            ms, job["result"]["generated_version_id"]
        )
        generated = json.loads(
            Path(
                next(
                    a["storage_path"]
                    for a in version["artifacts"]
                    if a["artifact_kind"] == "GENERATED_JSON"
                )
            ).read_text(encoding="utf-8")
        )
        rendered = generated["sections"][0]["paragraphs"][0]
        assert rendered.get("asset_ids") == paragraph["asset_ids"], (
            "Original material IDs must reach output"
        )
        assert [a["asset_id"] for a in rendered["assets"]] == paragraph["asset_ids"]
        assert captured[0]["assets"] == rendered["assets"]
        assert [a["sha256"] for a in rendered["assets"]] == [
            hashlib.sha256(raw).hexdigest() for _, _, raw in files
        ]
        block = next(
            b
            for b in version["blocks"]
            if b["kind"] == "paragraph" and b["object_id"] == paragraph["id"]
        )
        assert block["assets"] == rendered["assets"]
        pdf = Path(
            next(
                a["storage_path"]
                for a in version["artifacts"]
                if a["artifact_kind"] == "PDF"
            )
        )
        text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
        assert all(
            label in text
            for label in [
                "CSV_ROW_10",
                "XLSX_ROW_10",
                "MD_ROW_10",
                "SECOND_SHEET_ORIGINAL",
            ]
        )
        tex = Path(
            next(
                a["storage_path"]
                for a in version["artifacts"]
                if a["artifact_kind"] == "TEX"
            )
        ).read_text(encoding="utf-8")
        assert tex.count(r"\includegraphics") == 3
        geometry_path = tmp_path / "actual-material-layout.xhtml"
        layout = subprocess.run(
            [shutil.which("pdftotext"), "-bbox-layout", str(pdf), str(geometry_path)],
            capture_output=True,
            text=True,
            check=False,
        )
        assert layout.returncode == 0, layout.stderr
        geometry = ET.parse(geometry_path).getroot()
        namespace = {"x": "http://www.w3.org/1999/xhtml"}
        for page in geometry.findall(".//x:page", namespace):
            for word in page.findall(".//x:word", namespace):
                assert 40 <= float(word.get("xMin"))
                assert float(word.get("xMax")) <= float(page.get("width")) - 40
        for title, first_row in [
            ("original-csv.csv", "CSV_ROW_1"),
            ("sheet", "XLSX_ROW_1"),
            ("sheet", "SECOND_SHEET_ORIGINAL"),
            ("original-markdown.md", "MD_ROW_1"),
        ]:
            row_pages = [
                page
                for page in geometry.findall(".//x:page", namespace)
                if any(
                    word.text == first_row
                    for word in page.findall(".//x:word", namespace)
                )
            ]
            assert len(row_pages) == 1
            words = row_pages[0].findall(".//x:word", namespace)
            first_y = next(float(w.get("yMin")) for w in words if w.text == first_row)
            assert any(
                word.text == title and float(word.get("yMax")) < first_y
                for word in words
            ), (title, first_row)
        from fastapi.testclient import TestClient

        from rstv4.studio_api import create_app

        with TestClient(create_app(tmp_path)) as client:
            base = f"/api/manuscripts/{ms}/generated-versions/{version['generated_version_id']}/materials"
            for asset, (_, mime, original) in zip(
                rendered["assets"], files, strict=True
            ):
                response = client.get(f"{base}/{asset['asset_id']}/source")
                assert response.status_code == 200 and response.content == original
                assert response.headers["content-type"].split(";")[0] == mime
                assert response.headers["x-content-type-options"] == "nosniff"
                if asset.get("render_file"):
                    response = client.get(f"{base}/{asset['asset_id']}/render")
                    assert (
                        response.status_code == 200
                        and hashlib.sha256(response.content).hexdigest()
                        == asset["render_sha256"]
                    )
            assert client.get(f"{base}/foreign/source").status_code == 404
            assert (
                client.get(f"{base}/{assets[0]['asset_id']}/unsupported").status_code
                == 404
            )
            assert (
                client.get(
                    base.replace(ms, "foreign")
                    + "/"
                    + assets[0]["asset_id"]
                    + "/source"
                ).status_code
                == 404
            )
        assert service.get_workspace(ms)["document"] == doc
        source_bytes = Path(assets[0]["storage_path"]).read_bytes()
        Path(assets[0]["storage_path"]).write_bytes(b"X" + source_bytes[1:])
        try:
            assert service.get_generated_version(ms, version["generated_version_id"])
            assert (
                service.generated_material_content(
                    ms, version["generated_version_id"], assets[0]["asset_id"], "source"
                )[0]
                == source_bytes
            )
            with pytest.raises(DocumentValidationError):
                service.create_job(ms, "GENERATE", input_revision=1)
        finally:
            Path(assets[0]["storage_path"]).write_bytes(source_bytes)
        rendered_path = pdf.parent / rendered["assets"][0]["render_file"]
        frozen_bytes = rendered_path.read_bytes()
        rendered_path.write_bytes(b"X" + frozen_bytes[1:])
        try:
            with pytest.raises(RSTV4Error, match="integrity"):
                service.get_generated_version(ms, version["generated_version_id"])
        finally:
            rendered_path.write_bytes(frozen_bytes)
        service.set_stage(
            ms,
            "FINALIZATION_SUBMISSION",
            expected_revision=service.get_workspace(ms)["manuscript"]["revision"],
            reason="Software material delivery",
        )
        package_job = wait_for_job(
            service, service.create_job(ms, "SUBMISSION_PACKAGE", input_revision=1)
        )
        assert package_job["status"] == "SUCCEEDED", package_job
        package = service.list_submission_packages(ms)[0]
        raw, _, _ = service.submission_package_content(ms, package["package_id"])
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            for asset, (_, _, original) in zip(rendered["assets"], files, strict=True):
                assert archive.read(asset["source_file"]) == original
                if asset.get("render_file"):
                    assert (
                        hashlib.sha256(archive.read(asset["render_file"])).hexdigest()
                        == asset["render_sha256"]
                    )
            independent = Path(tempfile.mkdtemp(prefix="sc-material-package-"))
            assert independent.parent.resolve() == Path(tempfile.gettempdir()).resolve()
            archive.extractall(independent)
        compile_result = subprocess.run(
            [
                shutil.which("xelatex"),
                "-no-shell-escape",
                "-interaction=nonstopmode",
                "-halt-on-error",
                "clean-manuscript.tex",
            ],
            cwd=independent,
            capture_output=True,
            text=True,
            check=False,
        )
        assert compile_result.returncode == 0, (
            compile_result.stdout + compile_result.stderr
        )
        (tmp_path / "verified-material-receipt.json").write_text(
            json.dumps(
                {
                    "manuscript_id": ms,
                    "generated_version_id": version["generated_version_id"],
                    "job_id": job["job_id"],
                    "pdf": str(pdf),
                    "generated": generated,
                    "package_directory": package_job["result"]["directory"],
                    "independent_recompile_directory": str(independent),
                    "actual_remote_models": 0,
                    "scientific_acceptance": False,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
    finally:
        service._executor.shutdown(wait=True)


@pytest.mark.parametrize("defect", ["corrupt", "foreign", "missing"])
def test_unavailable_materials_reject_before_any_job_or_model_write(tmp_path, defect):
    service, _, ms = create_studio(tmp_path)
    doc = complete_document(service.get_workspace(ms))
    asset = service.store_asset(
        ms, "original.csv", "text/csv", b"item,value\noriginal,1\n"
    )
    doc["sections"][0]["paragraphs"][0]["asset_ids"] = [asset["asset_id"]]
    if defect == "foreign":
        other = service.create_manuscript(
            service.track_service.get_manuscript(ms)["track_id"],
            "Other software manuscript",
        )["manuscript"]["manuscript_id"]
        asset = service.store_asset(
            other, "other.csv", "text/csv", b"item,value\nother,2\n"
        )
        doc["sections"][0]["paragraphs"][0]["asset_ids"] = [asset["asset_id"]]
    service.save_document(
        ms, base_revision=0, operation_id="material-failure", document=doc
    )
    if defect == "corrupt":
        path = Path(asset["storage_path"])
        raw = path.read_bytes()
        path.write_bytes(b"X" + raw[1:])
    elif defect == "missing":
        with sqlite3.connect(service.db_path) as conn, conn:
            conn.execute(
                "UPDATE studio_assets SET storage_path=? WHERE asset_id=?",
                (str(tmp_path / "absent.csv"), asset["asset_id"]),
            )
    with sqlite3.connect(service.db_path) as conn:
        assert conn.execute("PRAGMA wal_checkpoint(TRUNCATE)").fetchone() == (0, 0, 0)
    before = service.db_path.read_bytes()
    with pytest.raises(DocumentValidationError) as rejection:
        service.create_job(ms, "GENERATE", input_revision=1)
    assert any(i["path"] == asset["asset_id"] for i in rejection.value.issues)
    assert service.db_path.read_bytes() == before
    with sqlite3.connect(service.db_path) as conn:
        assert conn.execute("SELECT count(*) FROM studio_jobs").fetchone()[0] == 0
        assert (
            conn.execute("SELECT count(*) FROM studio_model_calls").fetchone()[0] == 0
        )
    assert service.get_workspace(ms)["document"] == doc


@pytest.mark.parametrize(
    "svg",
    [
        b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><style>@import url("https://example.invalid/font.css");</style></svg>',
        b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><rect style="fill:url(file:///secret)"/></svg>',
        b'<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1000000000000"/>',
    ],
)
def test_svg_external_resources_and_unbounded_dimensions_fail_before_render(svg):
    with pytest.raises(ValueError):
        checked_svg(svg)


def test_xlsx_without_cached_formula_and_markdown_extra_prose_are_not_silently_lost():
    workbook = Workbook()
    workbook.active.append(["Original formula", "=1+2"])
    buf = BytesIO()
    workbook.save(buf)
    with pytest.raises(ValueError, match="cached"):
        representation(
            {
                "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "filename": "original.xlsx",
            },
            buf.getvalue(),
        )
    with pytest.raises(ValueError, match="extra prose"):
        representation(
            {"media_type": "text/markdown", "filename": "original.md"},
            b"Original footnote\n|A|B|\n|---|---|\n|1|2|",
        )


def test_original_16bit_png_values_are_not_saturated_by_rgb_conversion():
    image = Image.frombytes("I;16", (2, 1), b"\x00\x10\x00\x80")
    output = BytesIO()
    image.save(output, format="PNG")
    original = output.getvalue()
    shown, render = representation(
        {"media_type": "image/png", "filename": "original16.png"}, original
    )
    assert render == original and shown["width"] == 2 and shown["height"] == 1
    with Image.open(BytesIO(render)) as restored:
        assert list(restored.getdata()) == [4096, 32768]
